# -*- coding: utf-8 -*-
"""
Testes automatizados do pipeline de triagem (Fases 1, 2 e 3).

Execução:  .venv/bin/python -m pytest tests/ -v
"""

from pathlib import Path

import pytest

from licenciamento.agente_administrativo import AgenteAdministrativo
from licenciamento.agente_financeiro import AgenteFinanceiro
from licenciamento.auditor_tecnico import (AuditorTecnico, MetricasRFO,
                                           MetricasSondagem)
from licenciamento.parser_formulario import FormularioParser, normalizar as normalizar_texto

RAIZ = Path(__file__).resolve().parents[1]
EXEMPLOS = RAIZ / "exemplos"


# ==============================================================================
# FASE 1 - Parser
# ==============================================================================
def _parse(nome: str) -> dict:
    return FormularioParser(str(EXEMPLOS / nome)).gerar_json()


def test_parser_extrai_empreendedor_e_pleito_lor():
    dados = _parse("formulario_LOR_medio_alto.htm")
    assert dados["empreendedor"]["nome_razao_social"] == "Madeireira Vale do Sinos Ltda"
    assert dados["empreendedor"]["cpf_cnpj"] == "12.345.678/0001-90"
    assert dados["pleito"]["tipo_licenca"] == "LOR"
    assert dados["pleito"]["fases_componentes"] == ["LP", "LI", "LO"]
    assert dados["status_triagem"] == "liberado_triagem"


def test_parser_extrai_empreendimento_e_coordenadas_utm():
    dados = _parse("formulario_LOR_medio_alto.htm")
    emp = dados["empreendimento"]
    assert emp["porte"] == "Médio"
    assert emp["potencial_poluidor"] == "Alto"
    assert emp["codram"] == "05.412.1-3"
    # matrícula = SOMENTE o número (a célula traz 'Matrícula nº 12.345 do
    # Cartório de Registro de Imóveis de Campo Bom')
    assert emp["matricula_imovel"] == "12.345"
    assert emp["area_total_ha"] == pytest.approx(3.80)
    coords = emp["coordenadas"]
    assert coords["formato"] == "UTM"
    assert coords["datum"] == "SIRGAS 2000"
    assert coords["easting_m"] == pytest.approx(469512.25)
    assert coords["northing_m"] == pytest.approx(6714382.10)
    assert coords["fuso"] == 22


def test_parser_coordenadas_lat_long_e_areas_em_m2():
    dados = _parse("formulario_LP_pequeno_baixo.htm")
    emp = dados["empreendimento"]
    assert emp["porte"] == "Pequeno"
    assert emp["potencial_poluidor"] == "Baixo"
    assert emp["area_total_ha"] == pytest.approx(0.075)   # 750 m² -> ha
    assert emp["coordenadas"]["latitude"] == pytest.approx(-29.6745123)
    assert emp["coordenadas"]["longitude"] == pytest.approx(-50.8739812)


def test_art_ausente_bloqueia_triagem():
    dados = _parse("formulario_LP_sem_art.htm")
    assert dados["status_triagem"] == "bloqueado_sem_art"
    assert not dados["responsavel_tecnico"]["registro_art"]


def test_desduplicacao_lor_lp_li_lo():
    """LOR = LP + LI + LO com desduplicação de documentos repetidos entre fases."""
    dados = _parse("formulario_LOR_medio_alto.htm")
    docs = dados["documentos_exigidos"]
    estat = docs["estatisticas"]
    # 20 itens bruto (6 LP + 7 LI + 7 LO, incluindo 'Cópia da Licença Prévia'
    # da LI e 'Cópia da Licença de Instalação' da LO, que são DOCUMENTOS)
    assert estat["total_bruto"] == 20
    assert estat["total_deduplicado"] == 16
    assert len(estat["removidos"]) == 4
    # a 'Cópia da matrícula do imóvel' (repetida em LP/LI/LO) só aparece uma vez
    matriculas = [d for d in docs["lista_deduplicada"] if "matrícula do imóvel" in d.lower()]
    assert len(matriculas) == 1


def test_deduplicacao_por_similaridade():
    itens = ["Cópia da matrícula do imóvel", "Cópia da matrícula do imovel",
             "Plano de Gerenciamento de Resíduos Sólidos (PGRS)", "Planta de localização"]
    dedup, removidos = FormularioParser._deduplicar_documentos(itens)
    assert len(dedup) == 3
    assert len(removidos) == 1


# ==============================================================================
# FASE 2 - Agentes
# ==============================================================================
def test_administrativo_bloqueia_sem_art_e_cruza_checklist():
    dados = _parse("formulario_LOR_medio_alto.htm")
    resultado = AgenteAdministrativo().auditar(dados, [])
    assert resultado["status_geral"] == "BLOQUEADO" or resultado["bloqueios"] == []
    # sem nenhum anexo, todos os exigidos ficam pendentes
    assert resultado["resumo"]["total_pendentes"] == 16

    resultado2 = AgenteAdministrativo().auditar(
        dados, ["formulario_enquadramento_assinado.pdf", "copia_cpf_cnpj.pdf"])
    assert resultado2["resumo"]["total_ok"] >= 2
    assert resultado2["status_geral"] == "PENDENTE"


def test_financeiro_porte_minimo_oficial():
    """Valores oficiais do Mínimo: LP=LI=LO=52,20 URMs (todos os potenciais)."""
    fin = AgenteFinanceiro()
    for potencial in ("Baixo", "Médio", "Alto"):
        r = fin.calcular_taxa("LP", "Mínimo", potencial, "Padaria")
        assert r["total_urm"] == pytest.approx(52.20)


def test_financeiro_soma_lir_lor():
    """TABELA A oficial (Lei 4.439/2015) + somatório da Res. COMDEMA 003/2017."""
    fin = AgenteFinanceiro()
    r_lor = fin.calcular_taxa("LOR", "Médio", "Alto", "Indústria de madeira")
    assert r_lor["composicao_fases"] == {"LP": 1523.60, "LI": 1508.20, "LO": 1969.60}
    assert r_lor["total_urm"] == pytest.approx(5001.40)

    r_lir = fin.calcular_taxa("LIR", "Pequeno", "Médio", "Oficina mecânica")
    assert r_lir["composicao_fases"] == {"LP": 143.90, "LI": 245.30}
    assert r_lir["total_urm"] == pytest.approx(389.20)


def test_financeiro_excecao_erb():
    fin = AgenteFinanceiro()
    r = fin.calcular_taxa("LI", "Mínimo", "Baixo",
                          "Estação de Rádio Base (ERB) / Transmissão")
    assert r["grupo_atividade"] == "ERB"
    # TABELA B oficial: transmissão/retransmissão = 1.960,00 URM por fase
    # (divergência com a especificação original 612/714/510 documentada no config)
    assert r["composicao_fases"] == {"LI": 1960.00}
    assert r["total_urm"] == pytest.approx(1960.00)
    assert r["porte_ou_faixa"] == "FIXO"


def test_financeiro_excecao_faixas_hectares():
    """TABELAS C e D oficiais: Lavra 0-5 ha (única publicada) e Parcelamento."""
    fin = AgenteFinanceiro()
    lavra = fin.calcular_taxa("LI", "Pequeno", "Baixo",
                              "Extração mineral - Pedreira", area_ha=7.2)
    assert lavra["grupo_atividade"] == "LAVRA_MINERAL"
    assert lavra["porte_ou_faixa"] == "5 a 10 ha"
    # Tabela C publica só 0-5 ha (Médio): acima disso repete a linha + aviso
    assert lavra["total_urm"] == pytest.approx(1105.10)
    assert any("não publicada" in a for a in lavra.get("avisos", []))

    lote = fin.calcular_taxa("LO", "Grande", "Alto",
                             "Parcelamento do solo - Loteamento", area_ha=3.0)
    assert lote["porte_ou_faixa"] == "0 a 5 ha"
    assert lote["total_urm"] == pytest.approx(1386.12)

    lote12 = fin.calcular_taxa("LOR", "Grande", "Alto",
                               "Parcelamento do solo - Loteamento", area_ha=12.0)
    assert lote12["porte_ou_faixa"] == "10 a 20 ha"
    assert lote12["total_urm"] == pytest.approx(2294.76 + 2568.96 + 2568.96)


def test_financeiro_comercio_tabela_e():
    """TABELA E: Comércio em Geral (Baixo) por área construída."""
    fin = AgenteFinanceiro()
    pequena = fin.calcular_taxa("LP", "Pequeno", "Baixo",
                                "Comércio em geral - loja", area_m2=30.0)
    assert pequena["grupo_atividade"] == "COMERCIO"
    assert pequena["total_urm"] == pytest.approx(25.00)

    media = fin.calcular_taxa("LOR", "Pequeno", "Baixo",
                              "Comércio varejista", area_m2=150.0)
    assert media["porte_ou_faixa"] == "50 a 200 m2"
    assert media["total_urm"] == pytest.approx(52.20 * 3)

    grande = fin.calcular_taxa("LI", "Pequeno", "Baixo",
                               "Loja de departamentos", area_m2=320.0)
    assert grande["total_urm"] == pytest.approx(77.20)

    # comércio de potencial Alto NÃO fica na Tabela E (só publica Baixo)
    alto = fin.calcular_taxa("LP", "Médio", "Alto",
                             "Comércio de produtos químicos")
    assert alto["grupo_atividade"] == "GERAL"

    # sem área informada: fallback Tabela A com aviso
    sem_area = fin.calcular_taxa("LP", "Pequeno", "Baixo", "Comércio em geral")
    assert sem_area["grupo_atividade"] == "GERAL"
    assert any("TABELA A" in a for a in sem_area.get("avisos", []))


def test_financeiro_tabela_f_autorizacoes():
    """TABELA F: taxa única de autorizações por tipo/quantidade."""
    fin = AgenteFinanceiro()
    r = fin.calcular_taxa("AUTORIZACAO", "Pequeno", "Baixo",
                          "Supressão de árvores",
                          tipo_autorizacao="supressao_arvores",
                          quantidade_autorizacao=30)
    assert r["grupo_atividade"] == "TABELA_F"
    assert r["total_urm"] == pytest.approx(50.00)  # 21 a 50 árvores

    r2 = fin.calcular_taxa("AUTORIZACAO", "Pequeno", "Baixo",
                           "Movimentação de terras",
                           tipo_autorizacao="movimentacao_terras",
                           quantidade_autorizacao=200)
    assert r2["total_urm"] == pytest.approx(150.00)  # mais de 150 m³

    # sem quantidade: menor faixa provisória + aviso
    r3 = fin.calcular_taxa("AUTORIZACAO", "Pequeno", "Baixo",
                           "Descapoeiramento", tipo_autorizacao="descapoeiramento")
    assert r3["total_urm"] == pytest.approx(20.00)
    assert any("provisória" in a for a in r3.get("avisos", []))

    # autorização genérica (sem tipo): valor interno provisório + aviso
    r4 = fin.calcular_taxa("AUTORIZACAO", "Pequeno", "Baixo", "Autorização ambiental")
    assert r4["total_urm"] == pytest.approx(52.20)
    assert any("Tabela F" in a for a in r4.get("avisos", []))


def test_financeiro_licenca_unica():
    """Licença Única (requerimento oficial): somatório LP+LI+LO."""
    fin = AgenteFinanceiro()
    r = fin.calcular_taxa("LICENCA_UNICA", "Mínimo", "Baixo", "Padaria")
    assert r["total_urm"] == pytest.approx(52.20 * 3)
    assert "Licença Única" in r["regra_aplicada"]


# ==============================================================================
# FASE 3 - Auditor Técnico
# ==============================================================================
def test_pontos_sondagem_regra_da_area():
    """TRs oficiais: RSCC = 3 pontos até 1 ha; Parcelamento = 4 furos até 1 ha."""
    at = AuditorTecnico()
    # Aterro RSCC (TR 2026): 3 pontos até 1,0 ha + 1 por hectare ou fração
    assert at.calcular_pontos_sondagem_exigidos(0.8, "RSCC") == 3
    assert at.calcular_pontos_sondagem_exigidos(1.0, "RSCC") == 3
    assert at.calcular_pontos_sondagem_exigidos(1.1, "RSCC") == 4
    assert at.calcular_pontos_sondagem_exigidos(3.5, "RSCC") == 6
    # Meio Físico/Parcelamento (TR 2025): 4 furos até 1 ha + 1 por ha ou fração
    assert at.calcular_pontos_sondagem_exigidos(0.8, "PARCELAMENTO") == 4
    assert at.calcular_pontos_sondagem_exigidos(3.5, "PARCELAMENTO") == 7
    assert at.calcular_ensaios_permeabilidade_exigidos(1.0, "RSCC") == 2
    assert at.calcular_ensaios_permeabilidade_exigidos(1.0, "PARCELAMENTO") == 3


def test_sondagem_rscc_distancia_e_impermeabilizacao():
    """TR Aterro RSCC: d entre 1,0 e 1,5 m -> abaixo do mínimo + impermeabilização."""
    at = AuditorTecnico()
    metricas = MetricasSondagem(contexto="RSCC", profundidade_lencol_m=2.00,
                                cota_base_aterro_m=1.00, area_ha=1.0,
                                furos_informados=3, ensaios_permeabilidade_informados=2)
    r = at.validar_sondagem_aterramento("laudo.pdf", metricas)
    assert r.status.value == "PENDENTE"
    assert any("1,5" in item for item in r.itens_reprovados)
    assert any("IMPERMEABILIZAÇÃO" in item for item in r.itens_reprovados)
    assert r.metricas["distancia_vertical_m"] == pytest.approx(1.0)


def test_sondagem_rscc_distancia_vedada():
    """TR Aterro RSCC: d < 1,0 m é terminantemente proibido."""
    at = AuditorTecnico()
    metricas = MetricasSondagem(contexto="RSCC", profundidade_lencol_m=1.20,
                                cota_base_aterro_m=0.50, area_ha=1.0,
                                furos_informados=3, ensaios_permeabilidade_informados=2)
    r = at.validar_sondagem_aterramento("laudo.pdf", metricas)
    assert r.status.value == "PENDENTE"
    assert any("PROIBIDA" in item or "IMPOSSIBILIDADE" in item for item in r.itens_reprovados)


def test_sondagem_rscc_conforme():
    """d >= 2,0 m, sondagens e ensaios suficientes -> CONFORME."""
    at = AuditorTecnico()
    metricas = MetricasSondagem(contexto="RSCC", profundidade_lencol_m=3.50,
                                cota_base_aterro_m=0.50, area_ha=1.0,
                                furos_informados=3, ensaios_permeabilidade_informados=2)
    r = at.validar_sondagem_aterramento("laudo.pdf", metricas)
    assert r.status.value == "CONFORME"


def test_rfo_densidade_2000_gera_pendente():
    """Cenário da especificação: densidade de 2.000 mudas/ha -> PENDENTE."""
    at = AuditorTecnico()
    metricas = MetricasRFO(nativos_suprimidos=120, exoticos_suprimidos=30,
                           mudas_nativas_propostas=1900, mudas_exoticas_propostas=90,
                           densidade_proposta_mudas_ha=2000.0)
    r = at.validar_rfo("laudo_rfo.pdf", metricas)
    assert r.status.value == "PENDENTE"
    assert r.metricas["mudas_exigidas"] == 1890  # 15*120 + 3*30
    assert any("2.000" in item and "3.000" in item for item in r.itens_reprovados)


def test_rfo_conforme():
    at = AuditorTecnico()
    metricas = MetricasRFO(nativos_suprimidos=10, exoticos_suprimidos=0,
                           mudas_nativas_propostas=150, densidade_proposta_mudas_ha=3100.0,
                           monitoramento_anos=2)
    r = at.validar_rfo("laudo.pdf", metricas)
    assert r.status.value == "CONFORME"


def test_rfo_monitoramento_e_especies():
    """TR RFO: monitoramento >= 2 anos e espécies >= metade das suprimidas."""
    at = AuditorTecnico()
    metricas = MetricasRFO(nativos_suprimidos=10, exoticos_suprimidos=0,
                           mudas_nativas_propostas=150, densidade_proposta_mudas_ha=3100.0,
                           monitoramento_anos=1, especies_suprimidas=8, especies_plantadas=3)
    r = at.validar_rfo("laudo.pdf", metricas)
    assert r.status.value == "PENDENTE"
    assert any("monitoramento" in i.lower() and "2 anos" in i for i in r.itens_reprovados)
    assert any("espécies" in i.lower() for i in r.itens_reprovados)


def test_rfo_laudo_que_cita_minimo_normativo():
    """Laudo que transcreve o TR ('densidade mínima de 3.000 mudas/ha') não deve
    ter o mínimo confundido com a proposta (regressão)."""
    at = AuditorTecnico()
    texto = ("Medidas de reposição florestal devem considerar o plantio de 15 mudas com mais de "
             "um metro de altura por indivíduo nativo e 3 mudas com mais de um metro de altura por "
             "indivíduo exótico suprimido, conforme RESOLUÇÃO COMDEMA nº 02/2017. Os plantios devem "
             "ser feitos com uma densidade mínima de 3.000 mudas/hectare. "
             "Projeto com 120 indivíduos nativos e 30 exóticos suprimidos, plantio de 1.900 mudas "
             "nativas, densidade de 2.000 mudas/hectare, monitoramento por 2 anos.")
    metricas = at.extrair_parametros_rfo(texto)
    assert metricas.densidade_proposta_mudas_ha == pytest.approx(2000.0)  # não 3.000!
    r = at.validar_rfo("laudo_cita_tr.pdf", metricas)
    assert r.status.value == "PENDENTE"
    assert any("2.000" in i and "3.000" in i for i in r.itens_reprovados)


def test_gabarito_oficial_carregado():
    """config/gabarito_trs.json (TRs oficiais) é aplicado automaticamente."""
    at = AuditorTecnico()
    assert at.gabarito_revisado is True
    assert "campobom.rs.gov.br" in at.fonte_gabarito
    assert at.parametros["rscc_sondagem_base"] == 3            # TR Aterro RSCC 2026
    assert at.parametros["parcelamento_sondagem_base"] == 4    # TR Meio Físico 2025
    assert at.parametros["prad_monitoramento_minimo_anos"] == 2  # TR PRAD 5.7
    assert at.parametros["rfo_razao_minima_especies"] == 0.5   # TR RFO 3.3
    assert at.parametros["rscc_distancia_vedada_m"] == 1.0     # TR Aterro RSCC 2.4


def test_checklist_conteudo_eiv():
    """TR EIV: validação por checklist de conteúdo mínimo."""
    at = AuditorTecnico()
    texto_completo = ("ESTUDO DE IMPACTO DE VIZINHANÇA. Razão social: Empresa X - CNPJ 12.345/0001-90. "
                      "Logradouro: Rua A, Bairro Centro. Descrição do empreendimento: comércio. "
                      "Geração de tráfego e carga e descarga: ... Ruídos e vibrações: 60 decibéis. "
                      "Medidas de controle e medidas mitigadoras: ... ART anexada.")
    r_ok = at.validar_checklist_tr("eiv.pdf", "TR EIV", texto_completo,
                                   AuditorTecnico._checklists_padrao()["EIV"])
    assert r_ok.status.value == "CONFORME"

    texto_parcial = "Estudo de Impacto de Vizinhança. Razão social: Empresa X. CNPJ 12.345."
    r_parc = at.validar_checklist_tr("eiv2.pdf", "TR EIV", texto_parcial,
                                     AuditorTecnico._checklists_padrao()["EIV"])
    assert r_parc.status.value == "PENDENTE"
    assert r_parc.itens_reprovados  # lista os itens ausentes


def test_laudo_rfo_pdf_do_exemplo():
    texto = AuditorTecnico.extrair_texto(
        "laudo_rfo_densidade_2000.pdf",
        (EXEMPLOS / "laudos" / "laudo_rfo_densidade_2000.pdf").read_bytes())
    at = AuditorTecnico()
    resultado = at.validar_rfo("laudo_rfo_densidade_2000.pdf",
                               at.extrair_parametros_rfo(texto))
    assert resultado.status.value == "PENDENTE"
    assert resultado.itens_reprovados


# ==============================================================================
# MOTIVO DO ENCAMINHAMENTO À SEMA - leitura da MARCAÇÃO (dupla checagem)
# ==============================================================================
FORMULARIO_MOTIVO_TEMPLATE = """<!DOCTYPE html>
<html lang="pt-BR"><head><meta charset="utf-8"><title>Formulário</title></head>
<body>
<h1>SEMA CAMPO BOM</h1>
<table>
  <tr><td class="rotulo">Nome/Razão Social</td><td>Empresa Teste Ltda</td></tr>
  <tr><td class="rotulo">Nome do Empreendimento</td><td>Atividade Teste</td></tr>
</table>
<h2>3. MOTIVO DO ENCAMINHAMENTO À SEMA</h2>
<table>
  <tr><td>(  ) Licença Prévia (LP)</td></tr>
  <tr><td>(  ) Licença de Instalação (LI)</td></tr>
  <tr><td>(  ) Licença de Operação (LO)</td></tr>
  <tr><td>(  ) Licença de Instalação e Regularização (LIR)</td></tr>
  <tr><td>(  ) Licença de Operação e Regularização (LOR)</td></tr>
</table>
<h2>4. DOCUMENTAÇÃO EXIGIDA</h2>
<p><strong>Documentação exigida para a Licença Prévia (LP):</strong></p>
<ul>
  <li>Formulário assinado</li>
  <li>Cópia do CNPJ</li>
</ul>
<p><strong>Documentação exigida para a Licença de Instalação (LI):</strong></p>
<ul>
  <li>Cópia da Licença Prévia</li>
  <li>Projeto aprovado</li>
</ul>
<p><strong>Documentação exigida para a Licença de Operação (LO):</strong></p>
<ul>
  <li>PGRS</li>
</ul>
</body></html>"""


def test_formulario_oficial_marcado_lp_nao_sai_lor():
    """CENÁRIO CRÍTICO REPORTADO PELO LICENCIADOR: tabela oficial de opções em
    2 colunas com '( X ) Primeira licença' + '( X ) Licença Prévia' -> pleito
    LP (NUNCA LOR). Listagem cobrada = apenas a da LP; taxa cruza o
    Porte/Potencial COMBINADOS com o Manual (Tabela A: Pequeno/Baixo LP)."""
    dados = _parse("formulario_MOTIVO_LP_oficial.htm")
    pleito = dados["pleito"]
    assert pleito["tipo_licenca"] == "LP", pleito
    assert "marcação" in (pleito.get("metodo_deteccao") or "")
    assert pleito["fases_componentes"] == ["LP"]
    # listagem APENAS da fase marcada (5 documentos da LP; lista da LO fora)
    exigencias = dados["documentos_exigidos"]["lista_deduplicada"]
    assert len(exigencias) == 5
    assert any("matrícula do imóvel" in e.lower() for e in exigencias)
    assert not any("Bombeiros" in e for e in exigencias)

    emp = dados["empreendimento"]
    assert emp["porte"] == "Pequeno" and emp["potencial_poluidor"] == "Baixo"
    assert emp["area_intervencao_ha"] == pytest.approx(0.32)
    assert emp["codram"]  # CODRAM lido para o cruzamento com o Manual

    fin = AgenteFinanceiro().calcular_do_parser(dados)
    assert fin["total_urm"] == pytest.approx(72.10)
    assert fin["composicao_fases"] == {"LP": 72.10}


def test_motivo_encaminhamento_checkbox_lor():
    """Checkbox LOR marcado define o pleito; listagem = LP+LI+LO deduplicada."""
    dados = _parse("formulario_LOR_medio_alto.htm")
    pleito = dados["pleito"]
    assert pleito["tipo_licenca"] == "LOR"
    assert "marcação" in (pleito.get("metodo_deteccao") or "")
    assert pleito["fases_componentes"] == ["LP", "LI", "LO"]
    assert len(dados["documentos_exigidos"]["lista_deduplicada"]) == 16


def test_motivo_encaminhamento_marcacao_textual_lir():
    """'( X ) Licença de Instalação e Regularização (LIR)' -> LIR = LP + LI."""
    dados = _parse("formulario_MOTIVO_LIR.htm")
    pleito = dados["pleito"]
    assert pleito["tipo_licenca"] == "LIR"
    assert "marcação" in (pleito.get("metodo_deteccao") or "")
    assert pleito["fases_componentes"] == ["LP", "LI"]
    assert len(dados["documentos_exigidos"]["lista_deduplicada"]) == 7
    # item que menciona fase no TEXTO não pode ser engolido
    assert any("Licença Prévia" in d for d in
               dados["documentos_exigidos"]["lista_deduplicada"])


def test_pleitos_nomes_nas_duas_ordens():
    """As regularizações são reconhecidas nas DUAS ordens de nome usadas nos
    formulários, e a ordem das palavras não troca o tipo (LOR != LIR)."""
    padroes = dict(FormularioParser.PLEITOS)
    for texto in ["Licença de Operação e Regularização (LOR)",
                  "Licença de Regularização e Operação (LOR)",
                  "Licenca de Regularizacao e Operacao"]:
        assert padroes["LOR"].search(normalizar_texto(texto)), texto
        assert not padroes["LIR"].search(normalizar_texto(texto)), texto
    for texto in ["Licença de Instalação e Regularização (LIR)",
                  "Licença de Regularização e Instalação (LIR)",
                  "Licença de Implantação e Regularização (LIR)",
                  "Licenca de Regularizacao e Instalacao"]:
        assert padroes["LIR"].search(normalizar_texto(texto)), texto
        assert not padroes["LOR"].search(normalizar_texto(texto)), texto


def test_motivo_ordem_invertida_lir_e_lor():
    """Marcação com nomes na ordem invertida: LIR e LOR corretos + listagem."""
    rotulo_lir = "Licença de Regularização e Instalação (LIR)"
    html_lir = FORMULARIO_MOTIVO_TEMPLATE.replace(
        "(  ) Licença de Instalação e Regularização (LIR)", f"( X ) {rotulo_lir}")
    dados_lir = FormularioParser(conteudo_html=html_lir).parse()
    assert dados_lir["pleito"]["tipo_licenca"] == "LIR"
    assert dados_lir["pleito"]["fases_componentes"] == ["LP", "LI"]
    assert len(dados_lir["documentos_exigidos"]["lista_deduplicada"]) == 4

    rotulo_lor = "Licença de Regularização e Operação (LOR)"
    html_lor = FORMULARIO_MOTIVO_TEMPLATE.replace(
        "(  ) Licença de Operação e Regularização (LOR)", f"( X ) {rotulo_lor}")
    dados_lor = FormularioParser(conteudo_html=html_lor).parse()
    assert dados_lor["pleito"]["tipo_licenca"] == "LOR"
    assert dados_lor["pleito"]["fases_componentes"] == ["LP", "LI", "LO"]
    assert len(dados_lor["documentos_exigidos"]["lista_deduplicada"]) == 5


def test_secao_sem_marcacao_nao_adivinha_pleito():
    """Seção MOTIVO presente com TODAS as opções sem marca: NÃO adivinhar
    (a varredura geral casaria o LOR da lista) - fica None com aviso."""
    dados = FormularioParser(conteudo_html=FORMULARIO_MOTIVO_TEMPLATE).parse()
    assert dados["pleito"]["tipo_licenca"] is None
    avisos = " ".join(dados.get("avisos_parser") or [])
    assert "MARCADA" in avisos  # aviso de conferência manual para o licenciador


def test_marca_em_celula_isolada_da_tabela():
    """<td>X</td><td>Licença Prévia (LP)</td> (marca em célula própria)."""
    html = FORMULARIO_MOTIVO_TEMPLATE.replace(
        "<tr><td>(  ) Licença Prévia (LP)</td></tr>",
        "<tr><td>X</td><td>Licença Prévia (LP)</td></tr>")
    dados = FormularioParser(conteudo_html=html).parse()
    assert dados["pleito"]["tipo_licenca"] == "LP"


def test_nome_fantasia_no_metrico_do_frontend():
    """Empreendimento exibe o NOME FANTASIA; sem fantasia, a RAZÃO SOCIAL."""
    dados_fantasia = _parse("formulario_LOR_medio_alto.htm")
    assert dados_fantasia["empreendedor"].get("nome_fantasia") == "Serraria Vale do Sinos"

    dados_sem = _parse("formulario_MOTIVO_LIR.htm")
    assert not dados_sem["empreendedor"].get("nome_fantasia")
    fallback = (dados_sem["empreendedor"].get("nome_fantasia")
                or dados_sem["empreendedor"].get("nome_razao_social")
                or dados_sem["empreendimento"].get("nome_empreendimento"))
    assert fallback == "Oficina Mecânica Menezes Ltda"


def test_conferencia_arquivos_pela_listagem_do_motivo():
    """A conferência de anexos (Fase 2) usa a listagem gerada pela marcação."""
    dados = _parse("formulario_MOTIVO_LIR.htm")
    anexados = ["formulario_enquadramento_assinado.pdf", "copia_cpf_cnpj.pdf",
                "matricula_imovel.pdf", "art_responsavel_tecnico.pdf",
                "projeto_construcao_aprovado.pdf", "pgrs.pdf"]
    resultado = AgenteAdministrativo().auditar(dados, anexados)
    assert resultado["resumo"]["total_ok"] >= 5
    # o fixture LIR não tem ART (hard constraint -> BLOQUEADO é esperado);
    # o escopo deste teste é a CONFERÊNCIA DOS ANEXOS pela listagem marcada
    assert resultado["status_geral"] in ("APROVADO", "PENDENTE", "BLOQUEADO")


# ==============================================================================
# MOTIVO em DUAS COLUNAS (Primeira licença | Renovação) com marca em CÍRCULO
# ==============================================================================
def test_motivo_tabela_2col_circulo_primeira_licenca():
    """Formulário oficial: 1ª linha '[ X ] Primeira licença' | '[  ] Renovação'
    e tipos marcados com círculo '( o )'. O parser lê a PRIMEIRA LINHA, define
    a coluna ativa e depois o TIPO (LP); taxa cruza Porte/Potencial + CODRAM."""
    dados = _parse("formulario_MOTIVO_2col_circulo.htm")
    pleito = dados["pleito"]
    assert pleito["tipo_licenca"] == "LP", pleito
    assert pleito["natureza"] == "Primeira licença"
    assert "tabela" in (pleito.get("metodo_deteccao") or "")
    assert pleito["fases_componentes"] == ["LP"]
    # listagem SOMENTE da fase marcada (LP: 5 itens; LI/LO fora)
    docs = dados["documentos_exigidos"]["lista_deduplicada"]
    assert len(docs) == 5
    assert not any("Bombeiros" in d for d in docs)
    # dados lidos para o cruzamento com o Manual de Taxas
    emp = dados["empreendimento"]
    assert emp["porte"] == "Pequeno" and emp["potencial_poluidor"] == "Baixo"
    assert emp["area_intervencao_ha"] == pytest.approx(0.45)
    assert emp["codram"]
    fin = AgenteFinanceiro().calcular_do_parser(dados)
    assert fin["total_urm"] == pytest.approx(72.10)


def test_motivo_tabela_2col_circulo_renovacao():
    """Marcação na coluna RENOVAÇÃO: natureza 'Renovação' e o TIPO lido na
    coluna 2 (LI marcado) - a ordem das colunas não troca o tipo lido."""
    html = (EXEMPLOS / "formulario_MOTIVO_2col_circulo.htm").read_text(encoding="utf-8")
    html = html.replace("[ X ] <strong>Primeira licença</strong>",
                        "[&nbsp;&nbsp;&nbsp;] Primeira licença")
    html = html.replace("[&nbsp;&nbsp;&nbsp;] Renovação",
                        "[ X ] <strong>Renovação</strong>", 1)
    html = html.replace("""<td>( o ) Licença Prévia (LP)</td>
    <td>(&nbsp;&nbsp;&nbsp;) Licença Prévia (LP)</td>""",
                        """<td>(&nbsp;&nbsp;&nbsp;) Licença Prévia (LP)</td>
    <td>(&nbsp;&nbsp;&nbsp;) Licença Prévia (LP)</td>""")
    html = html.replace("""<td>(&nbsp;&nbsp;&nbsp;) Licença de Instalação (LI)</td>
    <td>(&nbsp;&nbsp;&nbsp;) Licença de Instalação (LI)</td>""",
                        """<td>(&nbsp;&nbsp;&nbsp;) Licença de Instalação (LI)</td>
    <td>( o ) Licença de Instalação (LI)</td>""")
    dados = FormularioParser(conteudo_html=html).gerar_json()
    pleito = dados["pleito"]
    assert pleito["tipo_licenca"] == "LI", pleito
    assert pleito["natureza"] == "Renovação"
    assert pleito["fases_componentes"] == ["LI"]


def test_matricula_validade_90_dias_corridos():
    """Validade da matrícula = 90 DIAS CORRIDOS da emissão (não 30)."""
    from datetime import date, timedelta
    from licenciamento.validador_documentos import ValidadorDocumentos
    validador = ValidadorDocumentos()
    assert validador.matricula_validade_dias == 90
    ref = date(2026, 9, 16)

    def texto_com_emissao(dias: int) -> str:
        emissao = ref - timedelta(days=dias)
        meses = {1: "janeiro", 2: "fevereiro", 3: "março", 4: "abril", 5: "maio",
                 6: "junho", 7: "julho", 8: "agosto", 9: "setembro", 10: "outubro",
                 11: "novembro", 12: "dezembro"}
        data_txt = f"{emissao.day} de {meses[emissao.month]} de {emissao.year}"
        return ("MATRÍCULA Nº 41.203 - Registro de Imóveis de Campo Bom/RS\n"
                "Imóvel: Rua Coronel João Corrêa, 215.\n"
                f"Campo Bom, {data_txt}.\n"
                "Oficial de Registro de Imóveis\n"
                "(documento assinado digitalmente)")

    dentro = validador.validar_matricula("matricula.txt", texto_com_emissao(89), ref)
    assert dentro.status.value == "CONFORME" or dentro.status.name == "CONFORME"
    vencida = validador.validar_matricula("matricula.txt", texto_com_emissao(91), ref)
    msg = " ".join(vencida.itens_reprovados)
    assert "90" in msg and "corridos" in msg
    # a norma citada também declara 'dias corridos'
    assert "corridos" in vencida.norma_tr


def test_sem_botao_exemplo_na_etapa_1():
    """Etapa 1 fica SOMENTE com o upload: sem botão 'Carregar exemplo fictício'."""
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file(str(RAIZ / "app.py"), default_timeout=120).run()
    assert not at.exception
    rotulos = [b.label for b in at.button]
    assert not any("exemplo" in r.lower() for r in rotulos), rotulos
    assert any("Analisar" in r for r in rotulos), rotulos


def test_painel_etapa2_2col_circulo_lp_e_taxa():
    """PONTE A PONTE do formulário do licenciador (2 colunas, círculo '( o )'):
    o painel da Etapa 2 mostra LP (não LOR) com taxa 72,10 URM e a natureza."""
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file(str(RAIZ / "app.py"), default_timeout=180)
    at.run()
    assert not at.exception
    at.radio[0].set_value("LP")                    # licença pleiteada
    at.radio[1].set_value("Primeira licença")      # natureza
    at.run()
    form_bytes = (EXEMPLOS / "formulario_MOTIVO_2col_circulo.htm").read_bytes()
    at.file_uploader[0].set_value(
        [("formulario_2col.htm", form_bytes, "text/html")])
    at.run()
    assert not at.exception
    [b for b in at.button if "Analisar" in b.label][0].click()
    at.run()
    assert not at.exception, [e.value[:300] for e in at.exception]
    metricas = {m.label: m.value for m in at.metric}
    assert metricas.get("Licença pleiteada") == "LP", metricas
    assert metricas.get("Taxa (URMs)") == "72,10", metricas
    captions = " | ".join(c.value for c in at.caption)
    assert "seleção sua na Etapa 1" in captions, captions
    assert "Primeira licença" in captions


# ==============================================================================
# Reconhecimento robusto do MOTIVO (spans do Word) + documentos com nome alterado
# ==============================================================================
def test_motivo_lista_simples_wingdings():
    """Formulário com a lista SIMPLES do item 3 (Licença Única, LP, LI, LO,
    LIR, LOR) e símbolos Wingdings do Word: vazio='o' em spans que o
    get_text separa do rótulo; marcado='þ' em Licença Prévia -> LP."""
    dados = _parse("formulario_MOTIVO_lista_simples.htm")
    pleito = dados["pleito"]
    assert pleito["tipo_licenca"] == "LP", pleito
    assert pleito["fases_componentes"] == ["LP"]
    assert len(dados["documentos_exigidos"]["lista_deduplicada"]) == 5
    fin = AgenteFinanceiro().calcular_do_parser(dados)
    assert fin["total_urm"] == pytest.approx(72.10)


def test_motivo_todos_com_circulo_menos_o_marcado():
    """Todas as opções com '( o )' e UMA com '( X )': a marca é o símbolo
    MINORITÁRIO (diferença entre as linhas) - e não a primeira linha."""
    linhas_opcoes = ["Licença Única", "Licença Prévia", "Licença de Instalação",
                     "Licença de Operação",
                     "Licença de Instalação e Regularização",
                     "Licença de Operação e Regularização"]
    corpo = "".join(
        f"<p>( {'X' if nome == 'Licença Prévia' else 'o'} ) {nome}</p>"
        for nome in linhas_opcoes)
    html = ("<!DOCTYPE html><html><body><h2>3. MOTIVO DO ENCAMINHAMENTO À SEMA</h2>"
            f"{corpo}<h2>4. DOCUMENTAÇÃO</h2><ul><li>Formulário assinado</li></ul>"
            "</body></html>")
    dados = FormularioParser(conteudo_html=html).gerar_json()
    assert dados["pleito"]["tipo_licenca"] == "LP", dados["pleito"]


def test_documento_nome_alterado_reconhecido_e_aprendido(tmp_path):
    """Pipeline do licenciador: nome alterado ('°Cópia da matrícula
    atualizada') reconhece PELO NOME; nome esdrúxulo ('doc_escaneado_0912')
    reconhece PELO CONTEÚDO e o sistema APRENDE o nome para os próximos
    processos; a matrícula ausente no formulário é COMPLETADA pelo anexo."""
    from licenciamento.identificador_documentos import IdentificadorDocumentos

    idf = IdentificadorDocumentos(caminho=tmp_path / "aprendido.json")
    html = (EXEMPLOS / "formulario_MOTIVO_lista_simples.htm").read_text(
        encoding="utf-8")
    dados = FormularioParser(conteudo_html=html).gerar_json()
    assert dados["empreendimento"].get("matricula_imovel") is None

    matricula_txt = ("CERTIDÃO DE INTEIRO TEOR\nMATRÍCULA Nº 41.203\n"
                     "Registro de Imóveis de Campo Bom/RS - Serventia e Registro\n"
                     "Campo Bom, 10 de setembro de 2026.")
    art_txt = ("Anotação de Responsabilidade Técnica\nART Nº 55210201145\n"
               "CREA-RS - Responsável Técnico: Eng. Ambiental Paulo Nunes")
    anexos = ["°Cópia da matrícula atualizada.txt", "doc_escaneado_0912.txt"]
    textos = {anexos[0]: matricula_txt, anexos[1]: art_txt}

    resultado = AgenteAdministrativo().auditar(
        dados, anexos, textos_anexados=textos, identificador=idf)
    # campo crítico completado pelo documento (não derruba mais o processo)
    assert dados["empreendimento"]["matricula_imovel"] == "41.203"
    assert any("COMPLETADO" in a for a in resultado["avisos"])
    assert not any("Matrícula" in b for b in resultado["bloqueios"])
    # ART esdrúxula: reconhecida pelo CONTEÚDO e APRENDIDA
    assert dados["responsavel_tecnico"]["registro_art"] == "55210201145"
    aprendidos = {i["nome"]: i["tipo"] for i in idf.aprendidos}
    assert aprendidos.get("doc escaneado 0912") == "ART"
    via_art = resultado["origem_ok"].get("ART do responsável técnico", {})
    assert via_art.get("via") in ("conteudo", "aprendido")

    # aprendizado persiste: nova instância reconhece a variação do nome
    idf2 = IdentificadorDocumentos(caminho=tmp_path / "aprendido.json")
    de_novo = idf2.identificar("doc_escaneado_0912_versao_final.txt", None)
    assert de_novo["tipo"] == "ART"
    assert de_novo["via"] in ("aprendido", "nome")


def test_painel_lista_simples_mostra_lp_e_taxa():
    """Ponte a ponte do formulário com a lista simples do item 3 (símbolos
    separados por spans): o painel exibe LP e a taxa 72,10 URM."""
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file(str(RAIZ / "app.py"), default_timeout=180)
    at.run()
    assert not at.exception
    at.radio[0].set_value("LP")
    at.radio[1].set_value("Primeira licença")
    at.run()
    form_bytes = (EXEMPLOS / "formulario_MOTIVO_lista_simples.htm").read_bytes()
    at.file_uploader[0].set_value(
        [("formulario_lista_simples.htm", form_bytes, "text/html")])
    at.run()
    assert not at.exception
    [b for b in at.button if "Analisar" in b.label][0].click()
    at.run()
    assert not at.exception, [e.value[:300] for e in at.exception]
    metricas = {m.label: m.value for m in at.metric}
    assert metricas.get("Licença pleiteada") == "LP", metricas
    assert metricas.get("Taxa (URMs)") == "72,10", metricas


# ==============================================================================
# OUTROS MEIOS de leitura: HTML bruto com inputs mapeados + confirmação manual
# ==============================================================================
def test_leitura_html_bruto_com_inputs_de_formulario():
    """MEIO ALTERNATIVO: quando get_text não entrega símbolo algum (controles
    <input> do Word), o parser reconstrói a seção do HTML BRUTO mapeando cada
    input para seu símbolo (marcado->'☒', vazio->'○') e reattacha ao rótulo."""
    html = ("<html><body><h2>3. MOTIVO DO ENCAMINHAMENTO À SEMA</h2><div>"
            "<input type=\"checkbox\"><span>Licença Única</span><br>"
            "<input type=\"checkbox\" checked><span>Licença Prévia</span><br>"
            "<input type=\"checkbox\"><span>Licença de Instalação</span><br>"
            "<input type=\"checkbox\"><span>Licença de Operação</span></div>"
            "<h2>4. DOCUMENTAÇÃO</h2></body></html>")
    parser = FormularioParser(conteudo_html=html)
    sinteticas = parser._linhas_da_secao_em_texto_sintetico()
    assert "☒ Licença Prévia" in sinteticas, sinteticas
    dados = parser.gerar_json()
    assert dados["pleito"]["tipo_licenca"] == "LP", dados["pleito"]


def test_selecao_etapa1_destrava_fluxo_sem_marca_no_formulario():
    """NOVA ABORDAGEM: o licenciador DECLARA o pleito na Etapa 1 (LP) e o
    formulário cuja marca não sobreviveu não trava nada - a listagem segue
    a LP e a taxa é calculada (72,10 URM, Tabela A Pequeno/Baixo)."""
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file(str(RAIZ / "app.py"), default_timeout=180)
    at.run()
    assert not at.exception
    at.radio[0].set_value("LP")                    # licença pleiteada
    at.radio[1].set_value("Primeira licença")      # natureza
    at.run()
    sem_marca = (EXEMPLOS / "formulario_MOTIVO_sem_marca.htm").read_bytes()
    at.file_uploader[0].set_value(
        [("formulario_sem_marca.htm", sem_marca, "text/html")])
    at.run()
    assert not at.exception
    [b for b in at.button if "Analisar" in b.label][0].click()
    at.run()
    assert not at.exception, [e.value[:300] for e in at.exception]
    metricas = {m.label: m.value for m in at.metric}
    assert metricas.get("Licença pleiteada") == "LP", metricas
    assert metricas.get("Taxa (URMs)") == "72,10", metricas
    captions = " | ".join(c.value for c in at.caption)
    assert "seleção sua na Etapa 1" in captions, captions


def test_divergencia_marcao_x_selecao_sinalizada():
    """Formulário MARCADO como LP, licenciador seleciona LOR: prevalece a
    seleção e a divergência é sinalizada no painel (nunca em silêncio)."""
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file(str(RAIZ / "app.py"), default_timeout=180)
    at.run()
    assert not at.exception
    at.radio[0].set_value("LOR")
    at.radio[1].set_value("Primeira licença")
    at.run()
    form_bytes = (EXEMPLOS / "formulario_MOTIVO_2col_circulo.htm").read_bytes()
    at.file_uploader[0].set_value(
        [("formulario_2col.htm", form_bytes, "text/html")])
    at.run()
    assert not at.exception
    [b for b in at.button if "Analisar" in b.label][0].click()
    at.run()
    assert not at.exception, [e.value[:300] for e in at.exception]
    metricas = {m.label: m.value for m in at.metric}
    assert metricas.get("Licença pleiteada") == "LOR", metricas
    avisos = [w.value for w in at.warning]
    assert any("DIVERGÊNCIA" in a for a in avisos), avisos


def test_parser_aplicar_pleito_manual_remonta_listagem():
    """aplicar_pleito_manual troca o tipo e RE-MONTA a listagem pelas fases
    (LI -> só a lista da LI; LOR -> LP+LI+LO do próprio formulário) e
    registra divergência quando a marcação lida era outra."""
    html = (EXEMPLOS / "formulario_MOTIVO_2col_circulo.htm").read_text(
        encoding="utf-8")
    parser = FormularioParser(conteudo_html=html)
    dados = parser.parse()
    assert dados["pleito"]["tipo_licenca"] == "LP"  # lido da marcação

    dados_li = parser.aplicar_pleito_manual("LI", "Renovação")
    assert dados_li["pleito"]["tipo_licenca"] == "LI"
    assert dados_li["pleito"]["fases_componentes"] == ["LI"]
    assert len(dados_li["documentos_exigidos"]["lista_deduplicada"]) == 3
    assert dados_li["pleito"]["natureza"] == "Renovação"

    qtde_li = len(dados_li["documentos_exigidos"]["lista_deduplicada"])
    dados_lor = parser.aplicar_pleito_manual("LOR", "Primeira licença")
    # aplicar_pleito_manual muta o MESMO dict (self.dados): a contagem do LI
    # precisa ser capturada antes
    assert dados_lor is dados_li
    assert dados_lor["pleito"]["fases_componentes"] == ["LP", "LI", "LO"]
    assert "DIVERGÊNCIA" in dados_lor["pleito"]["divergencia_selecao"]
    # listagem do LOR vem das 3 fases do PRÓPRIO formulário (deduplicada)
    assert len(dados_lor["documentos_exigidos"]["lista_deduplicada"]) > qtde_li


# ==============================================================================
# Parcelamento de solo (LP): 'Documentos Requeridos' numerados + RTs da seção 4.3
# ==============================================================================
def test_parcelamento_lp_lista_numerada_14_documentos():
    """Formulário de PARCELAMENTO (LP): a seção 'Documentos Requeridos' no
    final fornece a listagem CONSERVANDO a numeração (1..14); a taxa usa a
    Tabela D (faixa 10-20 ha) com a área lida do formulário."""
    dados = _parse("formulario_PARCELAMENTO_LP.htm")
    pleito = dados["pleito"]
    assert pleito["tipo_licenca"] == "LP", pleito
    lista = dados["documentos_exigidos"]["lista_deduplicada"]
    assert len(lista) == 14, lista
    assert lista[0].startswith("1. Diretrizes Urbanísticas")
    assert lista[3].startswith("4. Cópia da matrícula atualizada")
    assert lista[13].startswith("14. ART de profissional habilitado")
    emp = dados["empreendimento"]
    assert emp["area_total_ha"] == pytest.approx(12.5)
    assert emp["codram"]
    fin = AgenteFinanceiro().calcular_do_parser(dados)
    assert fin["total_urm"] == pytest.approx(2294.76)  # Tabela D, 10-20 ha, LP


def test_responsaveis_4_3_extraidos_e_cruzados_nos_anexos():
    """Seção 4.3 (demais responsáveis de etapas): RTT/ART de cada profissional
    é procurada nos documentos com NOME e REGISTRO batendo; ausente => aviso."""
    dados = _parse("formulario_PARCELAMENTO_LP.htm")
    rts = dados["responsaveis_etapas"]
    assert len(rts) == 2, rts
    assert rts[0]["nome"] == "Ana Prado Schneider"
    assert rts[0]["art_rtt"] == "RTT 55210098745"
    assert "CREA" in (rts[0]["registro"] or "")
    assert rts[1]["art_rtt"] == "RTT 55210112458"

    matricula_txt = ("MATRÍCULA Nº 55.888 Registro de Imóveis\n"
                     "Campo Bom, 10 de setembro de 2026.")
    laudo_txt = ("Laudo Geológico com Sondagem e Ensaios de Infiltração\n"
                 "Responsável Técnica: Ana Prado Schneider CREA 84512/D "
                 "RTT 55210098745")
    anexos = ["°Cópia da matrícula atualizada.txt", "copia_cnpj.pdf",
              "contrato_social.pdf", "laudo_geologico_sondagem.pdf"]
    textos = {anexos[0]: matricula_txt, anexos[3]: laudo_txt,
              anexos[1]: "Comprovante de Inscrição CNPJ 45.678.912/0001-03",
              anexos[2]: "Contrato Social da Vale Verde Incorporadora"}
    resultado = AgenteAdministrativo().auditar(
        dados, anexos, textos_anexados=textos)
    conf = resultado["conferencia_responsaveis"]
    assert conf[0]["encontrado"] is True
    assert conf[0]["anexo"] == "laudo_geologico_sondagem.pdf"
    assert "nome e registro" in conf[0]["nivel"]
    assert conf[1]["encontrado"] is False          # RTT do Carlos não veio
    assert any("55210112458" in a for a in resultado["avisos"])
    # documentos reconhecidos mesmo com o nº na frente do exigido
    assert any("matrícula" in d.lower() and "4." in d
               for d in resultado["documentos_ok"])


def test_deduplicacao_ignora_numeracao_da_listagem():
    """'1. Matrícula...' na LP e '2. Matrícula...' na LO são o MESMO documento:
    a desduplicação compara SEM a numeração e mantém a 1ª ocorrência (c/ nº)."""
    lista, removidos = FormularioParser._deduplicar_documentos(
        ["1. Matrícula do imóvel atualizada", "2. Matrícula do imóvel atualizada",
         "3. PGRS"])
    assert len(lista) == 2
    assert lista[0].startswith("1. Matrícula")
    assert any("2." in r["documento"] for r in removidos)


# ==============================================================================
# FORMULÁRIO REAL (Parcelamento/Condomínios - estrutura Maria Belle, Drive)
# ==============================================================================
REAL = "formulario_MARIA_BELLE_PARCELAMENTO.htm"


def test_formulario_real_leitura_completa_do_cabecalho():
    """Formulário real Word->HTML: Porte/Potencial ETIQUETADO ('Porte mínimo /
    Potencial Poluidor Médio'), CODRAM '3414,40 - ...', matrícula em célula
    vizinha ('Nº matrícula atual do imóvel'), áreas em m² e coordenadas
    'Lat.(º): -29.691629' SIRGAS 2000 decimal."""
    dados = _parse(REAL)
    emp = dados["empreendimento"]
    assert emp["porte"] == "Mínimo" and emp["potencial_poluidor"] == "Médio"
    assert emp["codram"] == "3414,40"
    assert emp["matricula_imovel"] == "33024"
    assert emp["area_intervencao_ha"] == pytest.approx(0.1784)  # 1.783,75 m²
    coords = emp["coordenadas"]
    assert coords["formato"] == "GEOGRAFICA"
    assert coords["latitude"] == pytest.approx(-29.691629)
    assert coords["longitude"] == pytest.approx(-50.054433)
    # ART lida do próprio formulário (seção 8)
    assert dados["responsavel_tecnico"]["registro_art"] == "202613404"


def test_formulario_real_listagem_por_tipo_de_licenca():
    """Listagens REAIS no final do formulário (parágrafos numerados, itens
    quebrados em várias linhas): LP=14, LI=14, LO=5; LOR=29 (33 com as
    repetições entre fases suprimidas: 'ART de profissional' e 'Arquivo
    KMZ/KML e DWG' aparecem 1x); RENOVAÇÃO tem listagem própria de 5 itens."""
    parser = FormularioParser(str(EXEMPLOS / REAL))
    parser.parse()

    lp = parser.aplicar_pleito_manual("LP", "Primeira licença")
    lista_lp = lp["documentos_exigidos"]["lista_deduplicada"]
    assert len(lista_lp) == 14, lista_lp
    assert lista_lp[0].startswith("1. Diretrizes Urbanísticas")
    assert lista_lp[7].startswith("8. Estudo de Impacto de Vizinhança")
    assert "com ART de responsável técnico habilitado;" in lista_lp[7]  # continuação
    assert lista_lp[13].startswith("14. ART de profissional habilitado")
    # item 14 LIMPO: o letreiro entre listagens ('Licença Prévia (LP), ...')
    # NÃO é colado no fim do item
    assert "Licença Prévia (LP)," not in lista_lp[13]

    li = parser.aplicar_pleito_manual("LI", "Primeira licença")
    assert len(li["documentos_exigidos"]["lista_deduplicada"]) == 14
    assert li["documentos_exigidos"]["lista_deduplicada"][0].startswith(
        "1. Projeto de instalação do abastecimento")

    lo = parser.aplicar_pleito_manual("LO", "Primeira licença")
    assert len(lo["documentos_exigidos"]["lista_deduplicada"]) == 5

    lor = parser.aplicar_pleito_manual("LOR", "Primeira licença")
    lista_lor = lor["documentos_exigidos"]["lista_deduplicada"]
    assert len(lista_lor) == 29  # 14 + 14 + 5 - 4 repetições suprimidas
    assert sum(1 for x in lista_lor if "ART de profissional habilitado" in x) == 1
    assert sum(1 for x in lista_lor if "KMZ/KML e DWG" in x) == 1

    ren = parser.aplicar_pleito_manual("LO", "Renovação")
    lista_ren = ren["documentos_exigidos"]["lista_deduplicada"]
    assert len(lista_ren) == 5
    assert lista_ren[0].startswith("1. Cópia da licença a ser renovada")


def test_formulario_real_taxa_tabela_d_parcelamento():
    """Taxa do parcelamento: TABELA D pela área de intervenção (1.783,75 m² =
    0,1784 ha -> faixa 0 a 5 ha) => LP 1.290,70 URM."""
    # fluxo real: o licenciador SELECIONA LP na Etapa 1 (formulário sem
    # marca legível na seção 3 -> o pleito vem da seleção)
    parser = FormularioParser(str(EXEMPLOS / REAL))
    dados = parser.aplicar_pleito_manual("LP", "Primeira licença")
    fin = AgenteFinanceiro().calcular_do_parser(dados)
    assert fin["total_urm"] == pytest.approx(1290.70)
    assert fin["composicao_fases"] == {"LP": 1290.70}


def test_formulario_real_rts_43_art_sem_prefixo_cruzados():
    """Seção 4.3 real: ART vem SEM prefixo na tabela ('17246618') - leitura
    posicional pelo cabeçalho; RTs cruzadas com os documentos apresentados."""
    dados = _parse(REAL)
    rts = dados["responsaveis_etapas"]
    nomes = [r["nome"] for r in rts]
    assert "Raquel Beckes" in nomes
    raquel = next(r for r in rts if r["nome"] == "Raquel Beckes")
    assert raquel["art_rtt"] == "ART 17246618"
    assert raquel["registro"] == "A67154-1"
    assert "urban" in (raquel["etapa"] or "").lower()
    assert sum(1 for r in rts if r["nome"]) == 4  # linhas vazias ignoradas

    texto_urbanistico = ("Projeto Urbanístico do Loteamento\n"
                         "Responsável Técnica: Raquel Beckes, Registro CAU A67154-1, "
                         "ART 17246618")
    anexos = ["projeto_urbanistico.pdf", "inventario_fauna.pdf"]
    textos = {anexos[0]: texto_urbanistico,
              anexos[1]: "Inventário de fauna - ART 99999999 de outro profissional"}
    resultado = AgenteAdministrativo().auditar(dados, anexos, textos_anexados=textos)
    conf = {c["profissional"]: c for c in resultado["conferencia_responsaveis"]}
    assert conf["Raquel Beckes"]["encontrado"] is True
    assert conf["Raquel Beckes"]["anexo"] == "projeto_urbanistico.pdf"
    assert conf["Keli Daiane Bernardes dos Santos"]["encontrado"] is False


def test_listagem_ancorada_no_titulo_documentos_requeridos():
    """A LISTAGEM é lida EXCLUSIVAMENTE após o título 'Documentos Requeridos'
    (no final do formulário): itens numerados do CORPO (Quadro diagnóstico
    '1. Existe banhado?') ficam FORA da exigência."""
    # fluxo real: licenciador seleciona LP na Etapa 1 (sem marca legível)
    parser = FormularioParser(str(EXEMPLOS / REAL))
    dados = parser.aplicar_pleito_manual("LP", "Primeira licença")
    lista = dados["documentos_exigidos"]["lista_deduplicada"]
    assert len(lista) == 14
    assert not any("banhado" in x.lower() for x in lista)
    assert not any("inundação" in x.lower() for x in lista)
    # fonte registrada como a listagem ancorada no título
    assert "Documentos Requeridos" in dados["documentos_exigidos"]["fonte_checklist"]


# ==============================================================================
# Etapa 1: 'Autorização Geral' e 'PRAD' como tipos selecionáveis
# ==============================================================================
def test_autorizacao_geral_na_etapa1_listagem_e_taxa():
    """'Autorização Geral' selecionada na Etapa 1: NÃO herda as listagens
    LP/LI/LO do formulário - usa o checklist oficial da espécie; taxa fixa
    provisória de 52,20 URM (com aviso)."""
    parser = FormularioParser(str(EXEMPLOS / REAL))
    parser.parse()
    d = parser.aplicar_pleito_manual("AUTORIZACAO", "Primeira licença")
    lista = d["documentos_exigidos"]["lista_deduplicada"]
    assert 3 <= len(lista) <= 12, lista
    assert not any("Diretrizes Urbanísticas" in x for x in lista)  # era da LP
    fin = AgenteFinanceiro().calcular_do_parser(d)
    assert fin["total_urm"] == pytest.approx(52.20)


def test_prad_na_etapa1_sem_listagem_e_taxa_sinalizada():
    """'PRAD - Plano de Recuperação de Área Degradada' na Etapa 1: tipo
    aceito (fases ['PRAD']); a listagem fica VAZIA até chegar o documento de
    referência (Drive) e a taxa é sinalizada explicitamente como não
    publicada no Manual (não é licença por fase)."""
    parser = FormularioParser(str(EXEMPLOS / REAL))
    parser.parse()
    d = parser.aplicar_pleito_manual("PRAD", "Primeira licença")
    assert d["pleito"]["tipo_licenca"] == "PRAD"
    assert d["pleito"]["fases_componentes"] == ["PRAD"]
    assert d["documentos_exigidos"]["lista_deduplicada"] == []
    fin = AgenteFinanceiro().calcular_do_parser(d)
    assert fin["total_urm"] is None
    assert "não publicada" in (fin.get("erro") or "")


def test_prad_e_autorizacao_no_painel_da_etapa2():
    """Ponte a ponte: seleção PRAD na Etapa 1 -> painel mostra 'PRAD' com o
    aviso da taxa; seleção AUTORIZACAO -> 'AUTORIZACAO' com 52,20 URM."""
    from streamlit.testing.v1 import AppTest
    for tipo, taxa_esperada in [("PRAD", "—"), ("AUTORIZACAO", "52,20")]:
        at = AppTest.from_file(str(RAIZ / "app.py"), default_timeout=180)
        at.run()
        assert not at.exception
        at.radio[0].set_value(tipo)
        at.radio[1].set_value("Primeira licença")
        at.run()
        fb = (EXEMPLOS / REAL).read_bytes()
        at.file_uploader[0].set_value([(f"{tipo}.html", fb, "text/html")])
        at.run()
        assert not at.exception
        [b for b in at.button if "Analisar" in b.label][0].click()
        at.run()
        assert not at.exception, [e.value[:300] for e in at.exception]
        metricas = {m.label: m.value for m in at.metric}
        assert metricas.get("Licença pleiteada") == tipo, metricas
        assert metricas.get("Taxa (URMs)") == taxa_esperada, metricas


# ==============================================================================
# BANCO DE CHECKLISTS POR TIPO DE FORMULÁRIO (cabeçalho do HTML + licença)
# ==============================================================================
def _formulario_do_tipo(header: str) -> str:
    return (f"<html><head><title>{header}</title></head><body>"
            f"<h1>Formulário para Licenciamento Ambiental de: {header}</h1>"
            "<h2>1. IDENTIFICAÇÃO</h2><table><tr><td>Nome/Razão Social:</td>"
            "<td>Teste Ltda</td></tr></table></body></html>")


def test_banco_cabecalho_comercios_e_servicos_lp():
    """Cabeçalho 'COMÉRCIOS E SERVIÇOS' + LP na Etapa 1: listagem vem do BANCO
    (16 docs), NÃO do corpo do formulário."""
    parser = FormularioParser(conteudo_html=_formulario_do_tipo("COMÉRCIOS E SERVIÇOS"))
    parser.parse()
    d = parser.aplicar_pleito_manual("LP", "Primeira licença")
    assert d["tipo_formulario"]["chave"] == "comercios_servicos"
    lista = d["documentos_exigidos"]["lista_deduplicada"]
    assert len(lista) == 16, lista
    assert lista[0].startswith("1. Formulário de Licenciamento")
    assert "banco de checklists" in d["documentos_exigidos"]["fonte_checklist"]


def test_banco_erb_lor_criacao_animais_lir_acude_unico():
    """ERB + LOR -> 13 docs próprios; Criação de Animais + LIR -> lista da LP
    (equivalência do documento de referência); Açude (processo único) -> 12
    docs para qualquer licença."""
    p_erb = FormularioParser(conteudo_html=_formulario_do_tipo("ESTAÇÃO RÁDIO-BASE"))
    p_erb.parse()
    d = p_erb.aplicar_pleito_manual("LOR", "Primeira licença")
    assert len(d["documentos_exigidos"]["lista_deduplicada"]) == 13
    assert any("Laudo radiométrico" in x for x in
               d["documentos_exigidos"]["lista_deduplicada"])

    p_cri = FormularioParser(conteudo_html=_formulario_do_tipo("CRIAÇÃO DE ANIMAIS"))
    p_cri.parse()
    d2 = p_cri.aplicar_pleito_manual("LIR", "Primeira licença")
    lista2 = d2["documentos_exigidos"]["lista_deduplicada"]
    assert len(lista2) == 11
    assert any("Criação Animal" in x for x in lista2)

    p_acude = FormularioParser(conteudo_html=_formulario_do_tipo("ABERTURA DE AÇUDE"))
    p_acude.parse()
    d3 = p_acude.aplicar_pleito_manual("AUTORIZACAO", "Primeira licença")
    assert len(d3["documentos_exigidos"]["lista_deduplicada"]) == 12
    assert any("Outorga" in x for x in d3["documentos_exigidos"]["lista_deduplicada"])


def test_banco_industriais_renovacao_e_lir_soma_dedup():
    """Industriais: RENOVAÇÃO tem listagem própria (12); LIR soma LP+LI do
    banco com desduplicação (menos que a soma bruta)."""
    p = FormularioParser(conteudo_html=_formulario_do_tipo("ATIVIDADES INDUSTRIAIS"))
    p.parse()
    d_ren = p.aplicar_pleito_manual("LP", "Renovação")
    assert len(d_ren["documentos_exigidos"]["lista_deduplicada"]) == 12
    assert "RENOVACAO" in d_ren["documentos_exigidos"]["por_fase"]

    d_lir = p.aplicar_pleito_manual("LIR", "Primeira licença")
    lista = d_lir["documentos_exigidos"]["lista_deduplicada"]
    assert 20 < len(lista) < 30  # 16 + 14 brutos, com repetidos suprimidos


def test_banco_nao_casa_condominios_nem_troca_listagem():
    """O formulário de CONDOMÍNIOS HORIZONTAIS/VERTICAIS (Maria Belle) NÃO é
    confundido com 'Parcelamento/Loteamentos' nem 'Desmembramento': sem
    entrada no banco, usa a listagem PRÓPRIA 'Documentos Requeridos' (14)."""
    parser = FormularioParser(str(EXEMPLOS / REAL))
    parser.parse()
    d = parser.aplicar_pleito_manual("LP", "Primeira licença")
    assert d["tipo_formulario"] is None
    lista = d["documentos_exigidos"]["lista_deduplicada"]
    assert len(lista) == 14
    assert "Documentos Requeridos" in d["documentos_exigidos"]["fonte_checklist"]


def test_codram_nao_arrasta_runon_da_tabela():
    """Bug reportado: campo CODRAM vinha com o código + TODO o texto seguinte
    do formulário (run-on da tabela), quebrando a coluna 'Tipo de
    empreendimento' da Etapa 2. Fica só o código."""
    html = ("<html><body><h1>LICENCIAMENTO AMBIENTAL</h1><table>"
            "<tr><td>Nome/Razão Social:</td><td>Empreendimento Teste</td></tr>"
            "<tr><td>Ramo de Atividade:</td><td>Parcelamento</td></tr>"
            "<tr><td>CODRAM:</td><td>3414,40 Parcelamento do solo para fins "
            "residenciais e mistos 2. IDENTIFICAÇÃO DO PLEITO "
            "3. IDENTIFICAÇÃO DO EMPREENDIMENTO Área Total (ha): 5,0 "
            "Matrícula do Imóvel: 12345</td></tr></table></body></html>")
    parser = FormularioParser(conteudo_html=html)
    parser.parse()
    assert parser.dados["empreendimento"]["codram"] == "3414,40"

    # formato oficial pontilhado com dígito verificador também se mantém
    html2 = html.replace("3414,40 Parcelamento", "05.412.1-3 Loteamento")
    parser2 = FormularioParser(conteudo_html=html2)
    parser2.parse()
    assert parser2.dados["empreendimento"]["codram"] == "05.412.1-3"


# ==============================================================================
# PDFs ESCANEADOS (OCR) + CONFERÊNCIAS CNPJ e ART/RTT com o formulário
# ==============================================================================
def _pdf_escaneado(texto: str) -> bytes:
    """Gera um PDF de IMAGEM (página escaneada simulada) com o texto pedido."""
    import io
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (1400, 400), "white")
    d = ImageDraw.Draw(img)
    try:
        fonte = ImageFont.truetype(
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 40)
    except Exception:
        fonte = None
    d.text((40, 160), texto, fill="black", font=fonte)
    buffer = io.BytesIO()
    img.save(buffer, "PDF", resolution=100)
    return buffer.getvalue()


def test_leitor_pdf_ocr_documento_escaneado():
    """Matrículas/ARTs chegam como PDF de imagens: o leitor detecta a ausência
    de camada de texto e roda OCR (texto vira legível p/ as conferências)."""
    from licenciamento.leitor_pdf import LeitorPDF
    if not LeitorPDF.ocr_disponivel():
        pytest.skip("OCR (rapidocr-onnxruntime) indisponível neste ambiente")
    import re as _re
    pdf = _pdf_escaneado("MATRICULA 33024 CNPJ 12.345.678/0001-95")
    texto, info = LeitorPDF.extrair(pdf)
    assert info["metodo"] == "ocr", (texto, info)
    assert "33024" in _re.sub(r"\D", "", texto), texto
    # plug completo: o validador também lê o escaneado
    from licenciamento.validador_documentos import ValidadorDocumentos
    v = ValidadorDocumentos()
    assert "33024" in _re.sub(r"\D", "",
                              v.extrair_texto("matricula.pdf", pdf))


def test_conferencia_cnpj_matricula_vs_formulario():
    """CNPJ do formulário x NÚMERO DE INSCRIÇÃO na matrícula (1ª linha/coluna):
    CONFERE, DIVERGENTE, NAO_ENCONTRADO, ANEXO_NAO_LEGIVEL e CPF->ignora."""
    from licenciamento.agente_administrativo import AgenteAdministrativo
    dados = {"empreendedor": {"cpf_cnpj": "12.345.678/0001-95"},
             "documentos_exigidos": {"lista_deduplicada": [
                 "Cópia da matrícula atualizada do imóvel"]}}

    def conferir(texto_matricula):
        textos = {"matricula_do_imovel.pdf": texto_matricula}
        return AgenteAdministrativo._conferir_cnpj_matricula(
            dados, ["matricula_do_imovel.pdf"], textos)

    # OCR cola tokens: regex tolerante reconhece mesmo assim
    r1 = conferir("Numero de Inscricao CNPJ12.345.678/0001-95MATRICULA 33024 FALHA 2")
    assert r1["status"] == "CONFERE"  # "2" solto não pode virar CNPJ
    # CNPJ com espaços quebrados pelo OCR
    r2 = conferir("N. INSCRICAO 12 345 678 0001 95 - matricula 33024")
    assert r2["status"] == "CONFERE"
    # divergente
    r = conferir("Inscricao CNPJ 98.765.432/0001-10 matricula 33024")
    assert r["status"] == "DIVERGENTE" and r["cnpj_encontrado"] == "98765432000110"
    # sem CNPJ no texto
    r3 = conferir("Matricula 33024 proprietario joao da silva")
    assert r3["status"] == "NAO_ENCONTRADO"
    # anexo sem texto (escaneado e OCR indisponível)
    r4 = conferir("")
    assert r4["status"] == "ANEXO_NAO_LEGIVEL"
    # formulário com CPF (11 dígitos): não há conferência de CNPJ
    dados_cpf = {**dados, "empreendedor": {"cpf_cnpj": "123.456.789-00"}}
    assert AgenteAdministrativo._conferir_cnpj_matricula(
        dados_cpf, ["matricula_do_imovel.pdf"],
        {"matricula_do_imovel.pdf": "CNPJ 12.345.678/0001-95"}) is None


def test_art_do_responsavel_principal_conferida():
    """TODAS as ARTs/RTTs são conferidas: o responsável principal (item 14)
    entra na conferência nº+nome, sem duplicar quando já listado na 4.3."""
    from licenciamento.agente_administrativo import AgenteAdministrativo
    dados = {"responsavel_tecnico": {"nome": "João Pedro Kessler",
                                     "registro_art": "14572403",
                                     "registro_crea": "RS233891"},
             "responsaveis_etapas": []}
    textos = {"art_joao.pdf": "ART 14572403 responsavel tecnico "
                              "JOAO PEDRO KESSLER registro RS233891"}
    conf = AgenteAdministrativo()._conferir_responsaveis_etapas(
        dados, ["art_joao.pdf"], textos)
    assert len(conf) == 1
    assert conf[0]["encontrado"] and "item 14" in conf[0]["etapa"]

    # mesma ART já listada na seção 4.3: NÃO duplica
    dados_dup = {**dados, "responsaveis_etapas": [
        {"nome": "João Pedro Kessler", "registro": "RS233891",
         "art_rtt": "14572403", "etapa": "Projeto"}]}
    conf2 = AgenteAdministrativo()._conferir_responsaveis_etapas(
        dados_dup, ["art_joao.pdf"], textos)
    assert len(conf2) == 1 and conf2[0]["etapa"] == "Projeto"


# ==============================================================================
# SEÇÕES 1 e 2 do formulário REAL (tabelas irregulares, vários pares por linha)
# ==============================================================================
def test_secoes_1_e_2_tabela_real_celula_a_celula():
    """Formulário real (tabelas quebradas 4x4/1x5/1x4/1x2, vários pares
    rótulo|valor por linha): cada campo extrai SOMENTE o seu valor - matrícula
    '33024' sem a área, endereço da seção 2 sem 'Bairro/CEP', endereço do
    EMPREENDEDOR (seção 1) distinto do do EMPREENDIMENTO (seção 2)."""
    dados = _parse("formulario_MARIA_BELLE_PARCELAMENTO.htm")
    emp = dados["empreendimento"]
    # matrícula: só o número (não junta 'Área total de intervenção')
    assert emp["matricula_imovel"] == "33024"
    # endereço do EMPREENDIMENTO (seção 2, após o CODRAM) - célula única
    assert emp["endereco"] == ("Esquina entre as RuasJosé Vargas e Alberto "
                               "Fleck, lote 27, quadra 03, s/n")
    assert "Bairro" not in (emp["endereco"] or "")
    # município limpo (seção 2 não tem Município -> usa o do empreendedor)
    assert emp["municipio"] == "Campo Bom"
    # coordenadas: valores certos e valores_brutos LIMITADOS (nunca o HTML todo)
    coord = emp["coordenadas"]
    assert coord["latitude"] == -29.691629 and coord["longitude"] == -50.054433
    brutos = coord.get("valores_brutos") or ""
    assert len(brutos) <= 200 and "<" not in brutos
    # campos da seção 2 intactos
    assert emp["porte"] == "Mínimo" and emp["potencial_poluidor"] == "Médio"
    assert emp["codram"] == "3414,40"
    # seção 1 (EMPREENDEDOR): nome, CNPJ e o ENDEREÇO PRÓPRIO (1ª ocorrência)
    assert dados["empreendedor"]["nome_razao_social"] == \
        "Maria Joaquina Empreendimentos Imobiliários LTDA"
    assert dados["empreendedor"]["cpf_cnpj"] == "43929749000120"
    parser = FormularioParser(str(EXEMPLOS / REAL))
    # seção 1 (EMPREENDEDOR): endereço/contato próprios, célula a célula
    ee = dados["empreendedor"]
    assert ee["endereco"] == "Av. Oscar Cirilo Ritzel"
    assert ee["numero"] == "220" and ee["bairro"] == "Centro"
    assert ee["cep"] == "93700-000" and ee["municipio"] == "Campo Bom"
    assert ee["telefone"] == "(51) 3598-2323"
    assert ee["email"] == "ermel@ermelcontabil.com.br"
    # áreas: intervenção != útil != total (célula '1.783,75m² (área útil
    # total: 3.245,49)' + Quadro de áreas 4.2 'Total 3245,49')
    assert emp["area_intervencao_ha"] == pytest.approx(0.1784)
    assert emp["area_util_ha"] == pytest.approx(0.3245)
    assert emp["area_total_ha"] == pytest.approx(0.3245)
    # nome do empreendimento: NÃO é o cabeçalho do quadro ('Metragem (m²)')
    assert emp["nome_empreendimento"] == \
        "Maria Joaquina Empreendimentos Imobiliários LTDA"
    # responsável técnico principal (seção 8): nome adjacente ao 'ART nº:'
    assert dados["responsavel_tecnico"]["nome"] == \
        "Keli Daiane Bernardes dos Santos"
    assert dados["responsavel_tecnico"]["registro_art"] == "202613404"


def test_valor_a_esquerda_do_rotulo():
    """Layout em que o valor vem na célula à ESQUERDA do rótulo (print do
    licenciador): 'Nº matrícula atual do imóvel' sem célula à direita pega o
    número da célula anterior."""
    html = ("<html><body><h1>LICENCIAMENTO AMBIENTAL</h1><table>"
            "<tr><td>33024</td><td>Nº matrícula atual do imóvel:</td></tr>"
            "</table></body></html>")
    parser = FormularioParser(conteudo_html=html)
    parser.parse()
    assert parser.dados["empreendimento"]["matricula_imovel"] == "33024"


# ==============================================================================
# ART/RTT COMO DOCUMENTO PRÓPRIO (conferência × formulário) + roteador de TRs
# ==============================================================================
def test_art_documento_conferida_com_formulario():
    """ART enviada como PDF (às vezes foto/scan -> OCR): NÃO recebe auditoria
    de laudo (TR fauna etc.); é conferida - NÚMERO + NOME - com as ARTs
    declaradas no formulário HTML (RT principal + seção 4.3). Vários ARTs por
    processo: cada documento é conferido individualmente."""
    from licenciamento.auditor_tecnico import AuditorTecnico
    auditor = AuditorTecnico()
    art_texto = ("ANOTAÇÃO DE RESPONSABILIDADE TÉCNICA - ART\n"
                 "Nº da ART: 14572403\n"
                 "Profissional: João Pedro Sandri Kessler\n"
                 "Registro CREA-RS: 233891\n"
                 "Objeto: Elaboração de laudo de fauna silvestre - "
                 "mastofauna, avifauna e herpetofauna na área do "
                 "empreendimento.\n")
    arts_form = [{"numero": "202613404", "nome": "Keli Daiane Bernardes dos Santos"},
                 {"numero": "14572403", "nome": "João Pedro Sandri Kessler"}]
    res = auditor.auditar_documento("Art*228_assinado.pdf", art_texto,
                                    arts_formulario=arts_form)
    assert len(res) == 1
    assert "Conferência ART" in res[0].norma_tr
    assert res[0].status.value == "CONFORME", res[0].itens_reprovados
    # ART de número NÃO declarado no formulário -> REVISAO_MANUAL com aviso
    art_outra = art_texto.replace("14572403", "99999999")
    res2 = auditor.auditar_documento("art_outro.pdf", art_outra,
                                     arts_formulario=arts_form)
    assert res2[0].status.value == "REVISAO_MANUAL"
    assert any("NÃO confere" in i for i in res2[0].itens_reprovados)
    # sem formulário carregado -> REVISAO_MANUAL explícito
    res3 = auditor.auditar_documento("art3.pdf", art_texto)
    assert res3[0].status.value == "REVISAO_MANUAL"


def test_laudo_fauna_continua_recebendo_tr_fauna():
    """Laudo de fauna DE VERDADE (sem padrão de ART) continua na auditoria do
    TR LFS; e um documento que só MENCIONA 'com ART de responsável técnico'
    NÃO é tratado como ART."""
    from licenciamento.auditor_tecnico import AuditorTecnico
    auditor = AuditorTecnico()
    assert AuditorTecnico.identificar_art_rtt(
        "LAUDO DE FAUNA SILVESTRE\nmastofauna avifauna herpetofauna\n"
        "elaborado de acordo com o TR desta secretaria, com ART de "
        "responsável técnico habilitado") is None
    laudo = ("LAUDO DE FAUNA SILVESTRE - LFS\n"
             "Metodologia: busca ativa com armadilhas de interceptação e "
             "queda (pitfall) para mastofauna e avifauna; busca passiva com "
             "recordação acustica. Campanhas em primavera e verão. "
             "Curva do coletor: suficientemente amostado.\n")
    res = auditor.auditar_documento("lfs.pdf", laudo)
    assert any("Fauna" in r.norma_tr for r in res)


def test_eiv_com_secao_de_fauna_nao_recebe_tr_fauna():
    """EIV (título declara) tem seção de fauna no corpo: o roteador aplica
    SOMENTE o TR do EIV - não pode aparecer PENDENTE de TR de fauna."""
    from licenciamento.auditor_tecnico import AuditorTecnico
    auditor = AuditorTecnico()
    eiv = ("ESTUDO DE IMPACTO DE VIZINHANÇA - EIV\n"
           "1. Identificação do empreendimento...\n"
           "5. Diagnóstico da área de influência: fauna local (mastofauna, "
           "avifauna) registrada no entorno; população e infraestrutura "
           "urbana; tráfego gerado; ventilação e sombreamento.\n")
    res = auditor.auditar_documento("eiv.pdf", eiv)
    normas = [r.norma_tr for r in res]
    assert any("Vizinhan" in n for n in normas), normas
    assert not any("Fauna" in n for n in normas), normas


def test_listagem_ignora_rodape_obs_multilinha():
    """O OBS. quebra em 2 linhas no formulário real: a 2ª linha ('adicionais
    ao processo...') NÃO vira item da listagem, e o letreiro do bloco
    seguinte não cola no item 14."""
    parser = FormularioParser(str(EXEMPLOS / REAL))
    parser.parse()
    dados = parser.aplicar_pleito_manual("LP", "Primeira licença")
    docs = dados["documentos_exigidos"]["lista_deduplicada"]
    assert len(docs) == 14
    juntado = [d for d in docs if "adicionais ao processo" in d.lower()
               or "condomínios horizontais" in d.lower()]
    assert not juntado, juntado
    assert docs[-1].startswith("14. ART de profissional")


def test_tema_escuro_toggle_na_barra_lateral():
    """A app oferece alternância de tema (claro/escuro) na barra lateral:
    existe o toggle, liga sem exceção e a página continua renderizando."""
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file(str(RAIZ / "app.py"), default_timeout=120)
    at.run()
    assert len(at.sidebar.toggle) == 1
    at.sidebar.toggle[0].set_value(True).run()
    assert not at.exception
    # desligar de volta também não pode quebrar
    at.sidebar.toggle[0].set_value(False).run()
    assert not at.exception


# ==============================================================================
# AGENTE AUDITOR DO SISTEMA (independente): detecta -> dupla-checa -> corrige
# ==============================================================================
def test_auditor_sistema_bateria_rapida_no_repositorio():
    """O agente roda a bateria rápida no repositório real: todo achado passa
    pela DUPLA CHECAGEM (confirmado) e NÃO sobra erro de severidade alta
    (ambiente, parser nas fixtures, pleitos/taxas, integração do app)."""
    from licenciamento.auditor_sistema import AuditorSistema
    auditor = AuditorSistema(com_testes=False)
    erros = auditor.auditar()
    for e in erros:
        assert e.confirmado, f"{e.id} não foi dupla-checado: {e.descricao}"
    altas = [e for e in erros if e.severidade == "alta"]
    assert not altas, [f"{e.id}: {e.descricao}" for e in altas]
    rel = auditor.relatorio()
    assert rel["resumo"]["total"] == len(erros)
    assert "erros" in rel and "log" in rel


def test_auditor_dupla_checagem_descarta_falso_positivo():
    """Achado que NÃO se repete na 2ª passada isolada é descartado como
    FALSO_POSITIVO (registrado no log), sem atacar correção."""
    from licenciamento.auditor_sistema import AuditorSistema, Erro

    class AuditorInstavel(AuditorSistema):
        def __init__(self):
            super().__init__(com_testes=False)
            self.chamadas = 0

        def verificar_higiene_arquivos(self):
            self.chamadas += 1
            if self.chamadas == 1:  # só na 1ª passada (instável)
                return [Erro(id="HIG-FANTASMA", severidade="baixa",
                             componente="x", descricao="achado instável",
                             evidencia="-", correcao="-")]
            return []

    auditor = AuditorInstavel()
    erros = auditor.auditar()
    assert not erros
    assert any("FALSO_POSITIVO" in l for l in auditor.log)


def test_auditor_corrige_erro_confirmado_e_revalida(tmp_path):
    """Erro CONFIRMADO com correção automática (newline ausente) é corrigido
    na hora e REVALIDADO: status vira CORRIGIDO e o arquivo fica íntegro."""
    from licenciamento.auditor_sistema import AuditorSistema
    cfg = tmp_path / "config"
    cfg.mkdir()
    (cfg / "exemplo.json").write_bytes(b'{"ok": true}')  # sem newline
    auditor = AuditorSistema(com_testes=False, raiz=tmp_path)
    erros = auditor.auditar()
    achados = [e for e in erros if e.id == "HIG-NEWLINE"]
    assert achados and achados[0].confirmado
    auditor.corrigir(instalar_pacotes=False)
    alvo = [e for e in auditor.erros if e.id == "HIG-NEWLINE"][0]
    assert alvo.status == "CORRIGIDO" and alvo.corrigido
    assert (cfg / "exemplo.json").read_bytes().endswith(b"\n")


def test_frontend_skill_frontend_design_assinatura():
    """Skill frontend-design aplicada: cabeçalho institucional + TRILHA DE
    ETAPAS (assinatura) visíveis nas duas telas, com a etapa corrente
    destacada, e legenda do semáforo no quadro da Etapa 2."""
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file(str(RAIZ / "app.py"), default_timeout=120)
    at.run()
    md_inicial = "\n".join(m.value for m in at.markdown)
    assert "sema-trilha" in md_inicial
    assert "SEMA · Campo Bom" in md_inicial
    at.radio[0].set_value("LP")
    at.radio[1].set_value("Primeira licença")
    at.run()
    fb = (EXEMPLOS / REAL).read_bytes()
    at.file_uploader[0].set_value([("condominios.html", fb, "text/html")])
    at.run()
    [b for b in at.button if "Analisar" in b.label][0].click()
    at.run()
    assert not at.exception
    md_analise = "\n".join(m.value for m in at.markdown)
    assert "Avaliar o dossiê" in md_analise
    legendas = "\n".join(c.value for c in at.caption)
    assert "Legenda: ✅ Em conformidade" in legendas


# ==============================================================================
# DUPLA CHECAGEM dos TRs (laudo geológico, LCV, fauna) ponto a ponto
# ==============================================================================
def test_trs_dupla_checagem_geologico_lcv_fauna():
    """Laudo geológico, cobertura vegetal e inventário de fauna são confrontados
    com os TRs do banco (gabarito_trs.json) SEMPRE em DUPLA CHECAGEM: 2ª
    execução idêntica + TR confirmado pelo título; inconformidades citadas
    ponto a ponto com a exigência do TR."""
    from licenciamento.auditor_tecnico import AuditorTecnico
    auditor = AuditorTecnico()

    # --- LAUDO GEOLÓGICO: entra por infiltração/geológico (sem 'sondagem') ---
    geo = ("LAUDO GEOLÓGICO DE CAMPO\nEnsaios de infiltração duplo anel e "
           "análise geotécnica do terreno. Lençol freático em 1,2 m.\n")
    res_geo = auditor.auditar_com_dupla_checagem("geo.pdf", geo)
    assert any("Meio Físico" in r.norma_tr for r in res_geo), \
        [r.norma_tr for r in res_geo]
    r0 = [r for r in res_geo if "Meio Físico" in r.norma_tr][0]
    assert r0.metricas["dupla_checagem"].startswith("OK")
    assert "tr_confirmado_pelo_titulo" in r0.metricas

    # --- LCV: cada inconformidade cita o ITEM do TR ---
    lcv = ("LAUDO DE COBERTURA VEGETAL\nDescrição do método de inventário "
           "florestal com esforço amostral.\n")
    res_lcv = auditor.auditar_com_dupla_checagem("lcv.pdf", lcv)
    r_lcv = [r for r in res_lcv if "Cobertura" in r.norma_tr][0]
    assert r_lcv.status.value == "PENDENTE"
    unidos = " ".join(r_lcv.itens_reprovados)
    assert "Item obrigatório" in unidos and "TR Laudo de Cobertura Vegetal" in unidos
    assert "fitossociológico" in unidos  # item específico do TR citado
    assert r_lcv.metricas["dupla_checagem"].startswith("OK")

    # --- FAUNA: metodologia do TR LFS ponto a ponto ---
    fauna = ("INVENTÁRIO DE FAUNA\nBusca ativa com armadilhas de interceptação "
             "e queda para mastofauna; amostragens em primavera.\n")
    res_f = auditor.auditar_com_dupla_checagem("fauna.pdf", fauna)
    r_f = [r for r in res_f if "Fauna" in r.norma_tr][0]
    unidos_f = " ".join(r_f.itens_reprovados)
    if r_f.status.value != "CONFORME":
        assert "TR LFS" in unidos_f or "TR Laudo de Fauna" in r_f.norma_tr


def test_trs_dupla_checagem_roteamento_divergente_e_instavel():
    """Se o documento se declara 'Laudo Geológico' mas nenhum TR correspondente
    foi aplicado, a dupla checagem de roteamento SINALIZA (nunca silencia);
    execuções instáveis (1ª x 2ª divergentes) também viram REVISAO_MANUAL."""
    from licenciamento.auditor_tecnico import (AuditorTecnico,
                                               OrigemAnalise,
                                               ResultadoValidacao)
    from licenciamento.esquemas_tecnicos import StatusValidacao
    auditor = AuditorTecnico()
    # título declara LCV, mas a disputa RFO x LCV mantém só o RFO:
    # a dupla checagem de roteamento SINALIZA a ausência do TR esperado
    texto = ("LAUDO DE COBERTURA VEGETAL\nTexto sobre reposição florestal: "
             "densidade de plantio de mudas nativas para compensação de "
             "indivíduos suprimidos.\n")
    res = auditor.auditar_com_dupla_checagem("misto.pdf", texto)
    assert any("roteamento" in r.norma_tr.lower() for r in res), \
        [r.norma_tr for r in res]
    # e o oposto: título geológico com sondagem aplicada NÃO sinaliza
    res_ok = auditor.auditar_com_dupla_checagem(
        "geo2.pdf", "LAUDO GEOLÓGICO\nSondagens e ensaios de infiltração; "
        "lençol freático a 2,5 m.\n")
    assert not any("roteamento" in r.norma_tr.lower() for r in res_ok)

    class AuditorInstavel(AuditorTecnico):
        def __init__(self):
            super().__init__()
            self.n = 0

        def validar_pca(self, nome_documento, texto):
            self.n += 1
            status = (StatusValidacao.CONFORME if self.n % 2
                      else StatusValidacao.REVISAO_MANUAL)
            return ResultadoValidacao(
                documento_analisado=nome_documento,
                norma_tr="TR PCA (2026, item 5.1)", status=status,
                itens_reprovados=[],
                metricas={"execucao": self.n},
                origem=OrigemAnalise.DETERMINISTICO)

    res2 = AuditorInstavel().auditar_com_dupla_checagem(
        "pca.pdf", "PLANO DE CONTROLE AMBIENTAL cronograma de relatórios "
                   "trimestrais supressão de vegetação")
    assert any("DIVERGENTES" in i for r in res2
               for i in r.itens_reprovados)


# ==============================================================================
# CORREÇÕES do licenciador: ART (nome nos 1os dados + atividade), projeto
# urbanístico (profissional + áreas), CNPJ 'Número de Inscrição', sem msg de TR
# ==============================================================================
def test_art_nome_nos_primeiros_dados_e_atividade_licenciamento():
    """O nome do profissional SEMPRE está nos primeiros dados da ART: cruza os
    tokens do formulário com o texto do documento (como faz com o número). A
    descrição da atividade contendo 'licenciamento ambiental' atribui a
    responsabilidade técnica ao RT da seção 8 do formulário."""
    from licenciamento.auditor_tecnico import AuditorTecnico
    auditor = AuditorTecnico()
    art = ("ART - Anotação de Responsabilidade Técnica\\nNº da ART: 202613404\\n"
           "Keli Daiane Bernardes dos Santos\\nCREA/CRBio 110544/03-D\\n"
           "Descrição sumária da atividade: LICENCIAMENTO AMBIENTAL do "
           "empreendimento.\\n")
    res = auditor.auditar_com_dupla_checagem(
        "Art*228_assinado.pdf", art,
        arts_formulario=[{"numero": "202613404",
                          "nome": "Keli Daiane Bernardes dos Santos",
                          "secao": "8"}])
    r = [x for x in res if "ART" in x.norma_tr][0]
    assert r.status.value == "CONFORME", r.itens_reprovados
    assert "RESPONSABILIDADE TÉCNICA" in (r.metricas or {}).get("papel", "")
    assert "licenciamento ambiental" in (r.trecho_referencia or "").lower()
    # sem a atividade de licenciamento: ART da seção 8 fica para verificar
    art2 = art.replace("LICENCIAMENTO AMBIENTAL do "
                       "empreendimento.", "laudo de fauna silvestre.")
    r2 = [x for x in auditor.auditar_com_dupla_checagem(
        "art2.pdf", art2,
        arts_formulario=[{"numero": "202613404",
                          "nome": "Keli Daiane Bernardes dos Santos",
                          "secao": "8"}]) if "ART" in x.norma_tr][0]
    assert r2.status.value == "REVISAO_MANUAL"
    assert any("licenciamento" in i.lower() for i in r2.itens_reprovados)


def test_projeto_urbanistico_profissional_e_areas_vs_formulario():
    """Projetos urbanísticos com plantas: dupla checagem SEMPRE - profissional
    que assina bate com o formulário E área total/útil bate com os valores do
    formulário (m² convertido para ha); divergência é citada ponto a ponto."""
    from licenciamento.auditor_tecnico import AuditorTecnico
    auditor = AuditorTecnico()
    arts = [{"numero": "202613404",
             "nome": "Keli Daiane Bernardes dos Santos", "secao": "8"}]
    areas = {"area_total_ha": 1.25, "area_util_ha": 0.3245}
    proj = ("PROJETO URBANÍSTICO - plantas de situação e quadro de áreas\\n"
            "Responsável técnico: Keli Daiane Bernardes dos Santos - "
            "ART 202613404\\nÁrea total: 12.500,00 m²\\n"
            "Área útil: 3.245,49 m²\\n")
    r = [x for x in auditor.auditar_com_dupla_checagem(
        "projeto.pdf", proj, arts_formulario=arts,
        areas_formulario=areas) if "urbanístico" in x.norma_tr][0]
    assert r.status.value == "CONFORME", r.itens_reprovados
    # área divergente (13.000 m² != 12.500 m²): citada com os dois valores
    proj_err = proj.replace("12.500,00", "13.000,00")
    r2 = [x for x in auditor.auditar_com_dupla_checagem(
        "projeto2.pdf", proj_err, arts_formulario=arts,
        areas_formulario=areas) if "urbanístico" in x.norma_tr][0]
    assert r2.status.value == "PENDENTE"
    assert any("DIVERGE do formulário" in i and "1.3000" in i.replace(",", ".")
               for i in r2.itens_reprovados), r2.itens_reprovados


def test_sem_mensagem_tr_nao_reconhecido_e_cnpj_rotulo():
    """(1) Documentos que não são laudos NÃO geram 'Nenhum Termo de Referência
    reconhecido' (o usuário não envia TR). (2) CNPJ da matrícula sob o rótulo
    'Número de Inscrição' no formato xx.xxx.xxx/xxxx-xx é reconhecido."""
    from licenciamento.auditor_tecnico import AuditorTecnico
    res = AuditorTecnico().auditar_com_dupla_checagem(
        "matricula.pdf", "MATRÍCULA Nº 33024 do Cartório de Registro de "
                         "Imóveis - Numero de Inscrição do CNPJ "
                         "12.345.678/0001-95.")
    assert res == [] or not any("Termo de Referência reconhecido" in i
                                for r in res for i in r.itens_reprovados)
    from licenciamento.agente_administrativo import AgenteAdministrativo
    dados = {"empreendedor": {"cpf_cnpj": "12.345.678/0001-95"},
             "documentos_exigidos": {"lista_deduplicada": [
                 "Cópia da matrícula atualizada do imóvel"]}}
    c = AgenteAdministrativo._conferir_cnpj_matricula(
        dados, ["matricula.pdf"],
        {"matricula.pdf": "Cartório\\nNúmero de Inscrição: 12.345.678/0001-95"
                          "\\nCNPJ do proprietário acima"})
    assert c["status"] == "CONFERE", c


def test_agente_conformidade_verificacoes_conferencias():
    """Os erros reportados pelo licenciador viraram verificações permanentes do
    AGENTE DE CONFORMIDADE (CONF): ART/projeto/CNPJ funcionais + mensagem de TR
    abolida - a bateria roda limpa no repositório."""
    from licenciamento.auditor_sistema import AuditorSistema
    auditor = AuditorSistema(com_testes=False)
    erros = auditor.auditar()
    conf = [e for e in erros if e.id.startswith("CONF-")]
    assert not conf, [(e.id, e.descricao) for e in conf]


def test_projeto_urbanistico_sem_profissional_reconhecido_nao_crasha():
    """REGRESSÃO do ValidationError reportado pelo licenciador
    ('trecho_referencia Input should be a valid string'): projeto urbanístico
    cujo profissional NÃO confere (ou sem áreas legíveis) deve gerar PENDENTE
    com trecho_referencia STRING - nunca None/crash no painel."""
    from licenciamento.auditor_tecnico import AuditorTecnico
    auditor = AuditorTecnico()
    proj = ("PROJETO URBANÍSTICO - plantas\\nAssinado por outrem.\\n"
            "Área total: 500,00 m²\\n")
    res = auditor.auditar_com_dupla_checagem(
        "projeto_outrem.pdf", proj,
        arts_formulario=[{"numero": "202613404",
                          "nome": "Keli Daiane Bernardes dos Santos",
                          "secao": "8"}],
        areas_formulario={"area_total_ha": 1.25, "area_util_ha": 0.32})
    r = [x for x in res if "urbanístico" in x.norma_tr][0]
    assert r.status.value == "PENDENTE"
    assert isinstance(r.trecho_referencia, str)
    assert any("NÃO confere" in i for i in r.itens_reprovados)
    assert any("DIVERGE do formulário" in i for i in r.itens_reprovados)
    # e o modelo tolera None explicitamente (defesa em profundidade)
    from licenciamento.esquemas_tecnicos import (ResultadoValidacao,
                                                 StatusValidacao)
    r_none = ResultadoValidacao(documento_analisado="x", norma_tr="y",
                                status=StatusValidacao.PENDENTE,
                                itens_reprovados=[], trecho_referencia=None)
    assert r_none.trecho_referencia is None


def test_emissao_parecer_tecnico_download_funciona():
    """REWORK do parecer (usuário: 'botão ainda não aciona download'):
    o botão agora ACIONA O COMPILADOR DE TEXTO — prévia EDITÁVEL em
    text_area antes de exportar — e a exportação oferece .docx E .pdf,
    ambos com LINK data-URI embutido no markdown (rota /media/ não
    atravessa o proxy do preview). Fluxo testado ponta a ponta:
    (1) clicar 'Gerar parecer' compila o texto com o cabeçalho oficial;
    (2) a prévia editável contém 'PARECER TÉCNICO Nº';
    (3) docx decodificado do link abre como Word válido;
    (4) pdf decodificado valida com pypdf; (5) cópias em saidas/."""
    import base64
    import io
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file(str(RAIZ / "app.py"), default_timeout=180)
    at.run()
    at.radio[0].set_value("LP")
    at.radio[1].set_value("Primeira licença")
    at.run()
    fb = (EXEMPLOS / REAL).read_bytes()
    at.file_uploader[0].set_value([("condominios.html", fb, "text/html")])
    at.run()
    [b for b in at.button if "Analisar" in b.label][0].click()
    at.run()
    assert not at.exception
    at.checkbox[0].check().run()
    assert not at.exception

    # Simula preenchimento manual do Nº do parecer pelo usuário (campo vem limpo por padrão)
    for ti in at.text_input:
        if "parecer" in ti.label.lower():
            ti.set_value("001/2026")
    at.run()

    # (1) o botão agora COMPILA o texto em vez de baixar direto
    gerar = [b for b in at.button if "Gerar parecer" in b.label]
    assert gerar, "botão 'Gerar parecer (texto editável)' ausente"
    gerar[0].click()
    at.run()
    assert not at.exception

    # (2) prévia EDITÁVEL presente e com o cabeçalho do parecer
    area = [ta for ta in at.text_area
            if "PARECER TÉCNICO" in (ta.value or "")]
    assert area, "prévia editável do parecer não apareceu"
    assert "6. CONCLUSÃO" in area[0].value

    # (3)(4) exportações .docx e .pdf com links data-URI decodificáveis
    md = "\n".join(m.value for m in at.markdown)
    assert md.count("base64,") >= 2, "faltou link data-URI (docx/pdf)"
    links = []
    pos = 0
    for _ in range(2):
        ini = md.index("base64,", pos) + len("base64,")
        fim_l = md.index('"', ini)
        links.append(base64.b64decode(md[ini:fim_l]))
        pos = fim_l
    from docx import Document
    doc = Document(io.BytesIO(links[0]))  # .docx abre como Word válido
    assert "PARECER TÉCNICO Nº 001/2026" in "\n".join(
        p.text for p in doc.paragraphs)
    assert links[1][:5] == b"%PDF-"
    from pypdf import PdfReader
    leitor = PdfReader(io.BytesIO(links[1]))
    assert "PARECER TÉCNICO" in "\n".join(
        (pg.extract_text() or "") for pg in leitor.pages)

    # (5) cópias auditáveis
    assert (RAIZ / "saidas" / "parecer_tecnico_001-2026.docx").exists()
    assert (RAIZ / "saidas" / "parecer_tecnico_001-2026.pdf").exists()


def test_compilador_parecer_texto_unico_fonte():
    """O COMPILADOR DE TEXTO é a fonte única: compilar_texto_parecer
    produz as 6 seções + pendências; exportar_docx/exportar_pdf convertem
    o TEXTO (editado ou não) em arquivos válidos — inclui edição simulada
    do analista preservada na exportação."""
    import io
    from datetime import date
    from licenciamento.compilador_parecer import (compilar_texto_parecer,
                                                  exportar_docx, exportar_pdf)
    texto = compilar_texto_parecer(
        dados_processo={
            "empreendedor": {"nome_razao_social": "Maria Belle",
                             "cpf_cnpj": "43.929.749/0001-20"},
            "empreendimento": {"nome_empreendimento": "Condomínio Maria "
                                                  "Belle",
                               "ramo_atividade": "3414,40",
                               "porte": "Mínimo",
                               "potencial_poluidor": "Médio"},
            "pleito": {"tipo_licenca": "LP",
                       "fases_componentes": ["LP"]}},
        quadro_documentos=[
            {"documento": "Estudo Ambiental", "situacao": "PENDENTE",
             "pendencias": ["mapa de vizinhança ausente"],
             "arquivo": "estudo.pdf"},
            {"documento": "Termo de Referência — PCA",
             "situacao": "CONFORME", "pendencias": [], "arquivo": "pca.pdf"}],
        resultado_admin={"bloqueios": ["CNPJ divergente do formulário"],
                         "documentos_pendentes": []},
        resultados_tecnicos=[],
        arquivos_recebidos=["estudo.pdf", "pca.pdf"],
        resumo_quadro={"CONFORME": 1, "PENDENTE": 1, "NAO_APRESENTADO": 0},
        comentarios_analista="Conferido com o formulário oficial.",
        numero_parecer="001/2026", prazo_dias=30,
        data_referencia=date(2026, 9, 17))
    for marcador in ("PARECER TÉCNICO Nº 001/2026", "1. IDENTIFICAÇÃO",
                     "2. DOCUMENTAÇÃO", "3. ANÁLISE", "4. PENDÊNCIAS "
                     "ADMINISTRATIVAS", "5. ANÁLISE TÉCNICA", "6. CONCLUSÃO",
                     "CNPJ divergente", "prazo de 30 dias"):
        assert marcador in texto, marcador
    # analista edita o texto → a edição vai para as exportações
    texto_editado = texto.replace("Conferido com o formulário oficial.",
                                  "EDIÇÃO DO ANALISTA: conferido e ok.")
    docx = exportar_docx(texto_editado, "001/2026")
    from docx import Document
    corpo = "\n".join(p.text for p in Document(io.BytesIO(docx)).paragraphs)
    assert "EDIÇÃO DO ANALISTA" in corpo
    pdf = exportar_pdf(texto_editado, "001/2026")
    from pypdf import PdfReader
    cont = "\n".join((pg.extract_text() or "")
                      for pg in PdfReader(io.BytesIO(pdf)).pages)
    assert "EDIÇÃO DO ANALISTA" in cont


def test_seguranca_primitivas_e_auditor_sec():
    """Skills security-audit + senior-security (pedido do usuário): as
    primitivas de licenciamento/seguranca.py sanitizam DADO NÃO CONFIÁVEL
    (texto/nome vindo de documentos enviados) e o auditor tem a camada
    SEC- rodando limpa no repositório e DETETANDO regressões."""
    from licenciamento.seguranca import (attr_html, md_seguro,
                                         nome_arquivo_seguro, sufixo_seguro)
    # A. XSS armazenado: HTML é escapado, aspas/controle removidos
    sujo = '<script>alert(1)</script> a\tb'
    limpo = md_seguro(sujo)
    assert "<script>" not in limpo and "&lt;script&gt;" in limpo
    assert "\t" not in limpo and "<" not in limpo
    # B. injeção de link/imagem markdown: [ ] escapados -> sem âncora viva
    link = md_seguro("[clique aqui](javascript:alert(1))")
    assert "\\[" in link and "\\]" in link
    # C. slug de nome de arquivo: sem traversal nem aspas
    nome = nome_arquivo_seguro('../../.git/config" onerror="alert(1)')
    assert "/" not in nome and "\\" not in nome and '"' not in nome
    assert nome == ".git_config" + "_onerror_" + '"alert(1)'.replace('"', "_") \
        or ("passwd" not in nome and "/" not in nome)
    # D. sufixo: apenas [a-z0-9] com ponto, default .html
    assert sufixo_seguro("<script>").startswith(".")
    assert all(c in ".abcdefghijklmnopqrstuvwxyz0123456789"
               for c in sufixo_seguro("..%2F..%2F.HtmL!!"))
    assert sufixo_seguro("!!!") == ".html"
    # E. atributo HTML: aspas escapadas
    assert "&quot;" in attr_html('x" onmouseover="alert(1)')
    # F. gravação da caixa-preta usa sufixo seguro (integração)
    import re as _re
    from app import salvar_entrada_real

    class _Form:
        name = 'requerimento.html><script>..%2Fevil" onerror="x'
        getvalue = staticmethod(lambda: b"<html>teste</html>")

    salvar_entrada_real([_Form()], {})
    gravados = sorted(
        (RAIZ / "entradas_reais").glob("*_formulario*"),
        key=lambda q: q.stat().st_mtime)
    alvo = gravados[-1]
    assert _re.fullmatch(r"\.[a-z0-9]{1,10}", alvo.suffix), alvo.suffix
    assert "/" not in alvo.name[len(str(RAIZ)) + 1:]
    # G. auditor: camada SEC limpa no repositório real
    from licenciamento.auditor_sistema import AuditorSistema
    assert AuditorSistema(com_testes=False).verificar_seguranca() == []
