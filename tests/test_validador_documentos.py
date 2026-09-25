# -*- coding: utf-8 -*-
"""Testes do ValidadorDocumentos (análise documental e quadro resumo)."""
from datetime import date

import pytest

from licenciamento.esquemas_tecnicos import StatusValidacao
from licenciamento.validador_documentos import (ValidadorDocumentos, normalizar)

MATRICULA_MODELO = """
MATRÍCULA Nº 39.715 - Registro de Imóveis da Comarca de Campo Bom/RS
Imóvel: estrada municipal RS-239, distrito sede, lote 4 da quadra F.
Área total: 12.000,00 m². Proprietária: Madeireira Vale do Sinos Ltda,
inscrita no CNPJ sob o nº 12.345.678/0001-90.
Ônus: hipoteca em favor do Banco Fictício S.A. (contrato 901/2025).
Servidão de passagem registrada à margem.
Campo Bom, {data_longo}.
Oficial Registrador - Registro de Imóveis
Documento eletrônico assinado conforme Lei 11.419/2006 (canto inferior esquerdo)
"""


@pytest.fixture
def validador() -> ValidadorDocumentos:
    return ValidadorDocumentos()


@pytest.fixture
def pasta_tmp(tmp_path):
    """Pasta temporária para arquivos de teste."""
    return tmp_path


# ==============================================================================
# MATRÍCULA - data de emissão no FINAL do documento + prazo de 90 dias
# ==============================================================================
def test_matricula_dentro_do_prazo_90_dias(validador):
    """Emitida há 40 dias (canto inferior esquerdo) => CONFORME."""
    emissao = date(2026, 8, 6)  # 40 dias antes da referência 15/09/2026
    texto = MATRICULA_MODELO.format(data_longo="06 de agosto de 2026")
    r = validador.validar_matricula("matricula_imovel.pdf", texto,
                                    data_referencia=date(2026, 9, 15))
    assert r.status == StatusValidacao.CONFORME
    assert r.metricas["data_emissao"] == "2026-08-06"
    assert r.metricas["dias_desde_emissao"] == 40
    assert r.metricas["prazo_validade_dias"] == 90
    assert not r.itens_reprovados


def test_matricula_vencida_alem_de_90_dias(validador):
    """Emitida há mais de 90 dias => PENDENTE com mensagem do prazo."""
    texto = MATRICULA_MODELO.format(data_longo="05 de março de 2026")
    r = validador.validar_matricula("matricula.pdf", texto,
                                    data_referencia=date(2026, 9, 15))
    assert r.status == StatusValidacao.PENDENTE
    assert any("90 dias" in item for item in r.itens_reprovados)
    assert r.metricas["dias_desde_emissao"] == 194
    assert r.metricas["dias_restantes"] == -104


def test_matricula_data_nao_localizada(validador):
    """Sem data no final => REVISAO_MANUAL (conferência humana)."""
    texto = ("MATRÍCULA Nº 39.715 - Registro de Imóveis. Área total 12.000 m². "
             "Proprietária Madeireira Ltda.")
    r = validador.validar_matricula("matricula.pdf", texto,
                                    data_referencia=date(2026, 9, 15))
    assert r.status == StatusValidacao.REVISAO_MANUAL
    assert any("localizada" in item.lower() for item in r.itens_reprovados)


def test_matricula_data_no_formato_curto(validador):
    """Formato dd/mm/aaaa no fechamento também é aceito."""
    texto = ("MATRÍCULA Nº 12.345. Registrada no Livro 2. "
             "Campo Bom/RS, 10/07/2026.\nOficial de Registro")
    r = validador.validar_matricula("matricula.pdf", texto,
                                    data_referencia=date(2026, 9, 15))
    assert r.status == StatusValidacao.CONFORME  # 67 dias <= 90
    assert r.metricas["data_emissao"] == "2026-07-10"


def test_matricula_prazo_calibravel(validador):
    """O prazo de validade vem de config/regras_documentos.json (90 dias)."""
    assert validador.matricula_validade_dias == 90
    assert validador.regras_revisadas is True
    assert any("90 DIAS" in nota.upper() for nota in validador.notas_regras)


# ==============================================================================
# IDENTIFICAÇÃO E ANÁLISE POR TIPO DE DOCUMENTO
# ==============================================================================
def test_identificar_matricula_por_conteudo(validador):
    texto = MATRICULA_MODELO.format(data_longo="06 de agosto de 2026")
    assert validador.identificar_tipo("documento_qualquer.pdf", texto) == \
        "MATRICULA_IMOVEL"


def test_identificar_por_nome_do_arquivo(validador):
    assert validador.identificar_tipo("matricula_imovel_atualizada.pdf", "") == \
        "MATRICULA_IMOVEL"
    assert validador.identificar_tipo("pgrs_da_empresa.pdf", "") == "PGRS"
    assert validador.identificar_tipo("nota_qualquer.pdf", "texto solto") is None


def test_analisar_documento_pdf_sem_texto(validador):
    """Anexo escaneado/vazio => REVISAO_MANUAL com orientação."""
    r = validador.analisar_documento("laudo_escaneado.pdf", "",
                                     data_referencia=date(2026, 9, 15))
    assert r.status == StatusValidacao.REVISAO_MANUAL
    assert any("sem texto" in i.lower() for i in r.itens_reprovados)


def test_analisar_matricula_integra(validador):
    """analisar_documento roteia matrícula para a validação do prazo."""
    texto = MATRICULA_MODELO.format(data_longo="05 de março de 2026")
    r = validador.analisar_documento("matricula.pdf", texto,
                                     data_referencia=date(2026, 9, 15))
    assert r.status == StatusValidacao.PENDENTE
    assert r.metricas.get("prazo_validade_dias") == 90


def test_normalizar_utilitaria():
    assert normalizar("MATRÍCULA do Imóvel  nº 12") == "matricula do imovel no 12"


# ==============================================================================
# QUADRO RESUMO (checklist x arquivos x análises)
# ==============================================================================
def test_quadro_resumo_completo(validador):
    """Cobra exigências não apresentadas e marca pendências por arquivo."""
    exigencias = ["Cópia da matrícula atualizada do imóvel",
                  "Cópia do CNPJ",
                  "Plano de Gerenciamento de Resíduos Sólidos",
                  "Laudo de Cobertura Vegetal (LCV)"]
    matricula_txt = MATRICULA_MODELO.format(data_longo="05 de março de 2026")
    arquivos = [
        {"nome": "matricula_imovel.pdf", "texto": matricula_txt, "tipo": "MATRICULA_IMOVEL"},
        {"nome": "copia_cnpj.pdf", "texto": "COMPROVANTE DE INSCRIÇÃO CNPJ 12.345.678/0001-90",
         "tipo": "CNPJ"},
    ]
    analises = {
        "matricula_imovel.pdf": validador.analisar_documento(
            "matricula_imovel.pdf", matricula_txt, data_referencia=date(2026, 9, 15)),
        "copia_cnpj.pdf": validador.analisar_documento(
            "copia_cnpj.pdf", arquivos[1]["texto"]),
    }
    linhas, extras = validador.montar_quadro(exigencias, arquivos, analises)
    resumo = validador.resumo_quadro(linhas)

    por_exigencia = {linha["documento"]: linha for linha in linhas}
    # matrícula casada mas VENCIDA => pendente com a justificativa
    linha_mat = por_exigencia["Cópia da matrícula atualizada do imóvel"]
    assert linha_mat["situacao"] == "PENDENTE"
    assert linha_mat["arquivo"] == "matricula_imovel.pdf"
    assert any("90 dias" in p for p in linha_mat["pendencias"])
    # CNPJ casado e legível => em conformidade
    assert por_exigencia["Cópia do CNPJ"]["situacao"] == "CONFORME"
    # PGRS e LCV sem arquivo => não apresentado
    assert por_exigencia["Plano de Gerenciamento de Resíduos Sólidos"]["situacao"] == \
        "NAO_APRESENTADO"
    assert por_exigencia["Laudo de Cobertura Vegetal (LCV)"]["situacao"] == \
        "NAO_APRESENTADO"

    assert resumo == {"CONFORME": 1, "PENDENTE": 1, "NAO_APRESENTADO": 2}
    assert extras == []  # todos os arquivos casaram com alguma exigência


def test_quadro_arquivo_extra_sem_exigencia(validador):
    """Arquivo que não casa com nenhuma exigência volta na lista de extras."""
    exigencias = ["Cópia do CNPJ"]
    arquivos = [{"nome": "copia_cnpj.pdf", "texto": "CNPJ 12.345.678/0001-90"},
                {"nome": "foto_da_area.jpg.txt", "texto": "anotações da vistoria"}]
    linhas, extras = validador.montar_quadro(exigencias, arquivos, {})
    assert len(linhas) == 1 and linhas[0]["situacao"] == "CONFORME"
    assert extras and extras[0]["nome"] == "foto_da_area.jpg.txt"


# ==============================================================================
# EXTRAÇÃO MULTI-FORMATO
# ==============================================================================
def test_extrair_texto_docx_e_xlsx(validador, pasta_tmp):
    """Extração de Word e Excel (formatos aceitos na página de upload)."""
    from docx import Document
    import openpyxl

    doc = Document()
    doc.add_paragraph("PLANO DE GERENCIAMENTO DE RESÍDUOS SÓLIDOS")
    doc.add_paragraph("Empresa: Madeireira Vale do Sinos Ltda")
    doc.save(pasta_tmp / "pgrs.docx")
    r_doc = validador.extrair_texto("pgrs.docx", (pasta_tmp / "pgrs.docx").read_bytes())
    assert "PLANO DE GERENCIAMENTO" in r_doc

    plan = openpyxl.Workbook()
    aba = plan.active
    aba.title = "Resíduos"
    aba.append(["Tipo", "Volume mensal"])
    aba.append(["Orgânico", "120 kg"])
    plan.save(pasta_tmp / "residuos.xlsx")
    r_xls = validador.extrair_texto("residuos.xlsx",
                                    (pasta_tmp / "residuos.xlsx").read_bytes())
    assert "# Aba: Resíduos" in r_xls and "120 kg" in r_xls


def test_extrair_texto_formato_nao_suportado(validador):
    """Extensão estranha ou binário: devolve vazio sem quebrar (log)."""
    assert validador.extrair_texto("arquivo.zip", b"PK\x03\x04...") == ""


# ==============================================================================
# DOCUMENTOS ENVIADOS COMO IMAGEM (png/jpg)
# ==============================================================================
def test_identificar_imagem_pelo_nome(validador):
    """Matrícula enviada como foto casa com a exigência pelo NOME do arquivo."""
    assert validador.identificar_tipo("matricula_imovel.jpg", "") == "MATRICULA_IMOVEL"
    assert validador.identificar_tipo("matricula-do-imovel.png", "") == "MATRICULA_IMOVEL"
    assert validador.identificar_tipo("pgrs_foto.jpeg", "") == "PGRS"


def test_analisar_imagem_sem_ocr_revisao_manual(validador):
    """Imagem sem OCR => REVISAO_MANUAL com orientação + flag de imagem."""
    r = validador.analisar_documento("matricula_imovel.jpg", "",
                                     data_referencia=date(2026, 9, 15))
    assert r.status == StatusValidacao.REVISAO_MANUAL
    assert any("imagem" in item.lower() for item in r.itens_reprovados)
    assert r.metricas.get("imagem") is True
    assert r.metricas.get("ocr_disponivel") is False


def test_extrair_texto_imagem_nao_quebra(validador):
    """Bytes de imagem inválidos ou válidos: extração nunca levanta exceção."""
    assert validador.extrair_texto("foto.png", b"\x89PNG\r\n conteudo falso") == ""


def test_extrair_texto_png_real_sem_ocr(validador):
    """PNG legítimo (gerado com Pillow): sem tesseract devolve vazio (sem quebrar)."""
    from PIL import Image
    import io as _io
    buffer = _io.BytesIO()
    Image.new("RGB", (60, 30), color="white").save(buffer, format="PNG")
    texto = validador.extrair_texto("foto_area.png", buffer.getvalue())
    assert texto == ""


def test_extensoes_imagem_constante():
    """Formatos de imagem aceitos na página de upload."""
    for ext in (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif"):
        assert ext in __import__("licenciamento.validador_documentos",
                                 fromlist=["EXTENSOES_IMAGEM"]).EXTENSOES_IMAGEM


def test_quadro_casa_exigencia_com_imagem_bem_nomeada(validador):
    """Imagem 'matricula_imovel.jpg' atende a exigência da matrícula no quadro."""
    exigencias = ["Cópia da matrícula atualizada do imóvel"]
    arquivos = [{"nome": "matricula_imovel.jpg", "texto": "", "tipo": "MATRICULA_IMOVEL"}]
    analises = {"matricula_imovel.jpg": validador.analisar_documento(
        "matricula_imovel.jpg", "")}
    linhas, extras = validador.montar_quadro(exigencias, arquivos, analises)
    assert linhas[0]["arquivo"] == "matricula_imovel.jpg"
    assert linhas[0]["situacao"] == "PENDENTE"  # conferência manual pendente
    assert any("imagem" in p_.lower() for p_ in linhas[0]["pendencias"])
    assert extras == []


# ==============================================================================
# Casamento exigência x arquivo: NÚCLEO discriminante (sem cruzar documentos)
# ==============================================================================
def test_casamento_nao_cruza_documentos_casos_do_licenciador():
    """Casos reportados pelo licenciador: 'Diretrizes Urbanísticas' NÃO é dado
    como conforme pelo LAUDO DE COBERTURA VEGETAL; 'Inventário de fauna' não
    é satisfeito pelo EIV; 'Planta de localização' não é satisfeita pelo EIV."""
    validador = ValidadorDocumentos()
    eiv = {"nome": "eiv_estudo_impacto_vizinhanca.pdf",
           "texto": "Estudo de Impacto de Vizinhança (EIV), elaborado de acordo "
                    "com o TR desta secretaria, com ART de responsável técnico "
                    "habilitado. Contém planta de localização da área.",
           "tipo": "EIV"}
    laudo_lc = {"nome": "laudo_cobertura_vegetal.pdf",
                "texto": "Laudo de Cobertura Vegetal, elaborado de acordo com o "
                         "TR desta secretaria, com ART de responsável técnico "
                         "habilitado e indicação das APP.",
                "tipo": "LCV"}
    fauna = {"nome": "inventario_fauna.pdf",
             "texto": "Inventário de fauna, elaborado de acordo com o TR desta "
                      "secretaria, com ART de responsável técnico habilitado.",
             "tipo": None}
    arquivos = [eiv, laudo_lc, fauna]

    # Diretrizes Urbanísticas: NENHUM arquivo tem o núcleo -> NÃO_APRESENTADO
    assert validador.casar_exigencia(
        "1. Diretrizes Urbanísticas do Departamento de Planejamento desta "
        "Prefeitura;", arquivos) is None
    # EIV casa SOMENTE com o EIV (não com fauna, não com laudo de cobertura)
    assert validador.casar_exigencia(
        "8. Estudo de Impacto de Vizinhança (EIV), elaborado de acordo com o "
        "TR desta secretaria, com ART de responsável técnico habilitado;",
        arquivos) == "eiv_estudo_impacto_vizinhanca.pdf"
    # Inventário de fauna casa SOMENTE com o inventário (não com o EIV)
    assert validador.casar_exigencia(
        "9. Inventário de fauna, elaborado de acordo com o TR desta secretaria, "
        "com ART de responsável técnico habilitado;",
        arquivos) == "inventario_fauna.pdf"
    # Planta de localização (texto do EIV cita!) NÃO pode casar com o EIV
    # se existir o arquivo próprio
    planta = {"nome": "planta_localizacao.pdf",
              "texto": "Planta de localização da área com o entorno de 100 metros.",
              "tipo": None}
    arquivos_planta = arquivos + [planta]
    assert validador.casar_exigencia(
        "12. Planta de localização da área apresentando inclusive seu entorno "
        "(100 metros): vias de acesso, localização de recursos hídricos e APPs;",
        arquivos_planta) == "planta_localizacao.pdf"
    # sem o arquivo próprio: NÃO cai no EIV (núcleo 'planta localizacao'
    # presente só como menção) - fica NÃO apresentado
    assert validador.casar_exigencia(
        "12. Planta de localização da área apresentando inclusive seu entorno "
        "(100 metros): vias de acesso, localização de recursos hídricos e APPs;",
        arquivos) is None


def test_formulario_html_so_atende_exigencia_de_formulario():
    """O formulário .htm/.html NÃO é 'documento apresentado' para nenhuma
    outra condicionante - apenas para a exigência do próprio formulário."""
    validador = ValidadorDocumentos()
    formulario = {"nome": "formulario_MARIA_BELLE.html",
                  "texto": "Contrato social CNPJ matrícula atualizada "
                           "laudo de cobertura vegetal inventário de fauna "
                           "planta de localização diretrizes urbanísticas",
                  "tipo": "FORMULARIO"}
    assert validador.casar_exigencia(
        "4. Cópia da matrícula atualizada (últimos 90 dias);",
        [formulario]) is None
    assert validador.casar_exigencia(
        "2. Contrato social;", [formulario]) is None
    # exigência do próprio formulário: atende normalmente
    assert validador.casar_exigencia(
        "3. Formulário de Informações para Licenciamento Ambiental;",
        [formulario]) == "formulario_MARIA_BELLE.html"


def test_casamento_tolerante_a_nome_resumido():
    """Nomes de arquivo RESUMIDOS continuam casando: 'matricula_imovel.jpg'
    atende 'Cópia da matrícula atualizada do imóvel' (núcleo: 1ª palavra +
    apoio)."""
    validador = ValidadorDocumentos()
    imagem = {"nome": "matricula_imovel.jpg", "texto": "", "tipo": None}
    assert validador.casar_exigencia(
        "Cópia da matrícula atualizada do imóvel", [imagem]) == \
        "matricula_imovel.jpg"


def test_leitor_imagem_extracao_ocr():
    """Verifica a extração OCR via LeitorImagem e ValidadorDocumentos em imagem sintética."""
    import io
    from PIL import Image, ImageDraw
    from licenciamento.leitor_imagem import LeitorImagem
    from licenciamento.validador_documentos import ValidadorDocumentos

    img = Image.new("RGB", (500, 120), color=(255, 255, 255))
    d = ImageDraw.Draw(img)
    d.text((20, 40), "COMPROVANTE DE CNPJ 12.345.678/0001-90", fill=(0, 0, 0))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    conteudo = buf.getvalue()

    texto, info = LeitorImagem.extrair_texto(conteudo)
    assert "12.345.678/0001-90" in texto or "CNPJ" in texto
    assert info["sucesso"] is True

    validador = ValidadorDocumentos()
    texto_val = validador.extrair_texto("cartao_cnpj.png", conteudo)
    assert "12.345.678/0001-90" in texto_val or "CNPJ" in texto_val


# ==============================================================================
# TESTES DAS HABILIDADES GABEBRAIN NO LEITOR DE PDF E AUDITOR TÉCNICO
# ==============================================================================
def test_leitor_pdf_marcadores_pagina_e_fatiamento():
    """GabeBrain 02 (biblioteca-pesquisavel): extração por página [[pag N]],
    fatiamento seletivo e localização de página."""
    import pymupdf
    from licenciamento.leitor_pdf import LeitorPDF

    doc = pymupdf.open()
    pag1 = doc.new_page()
    pag1.insert_text((50, 50), "Capítulo 1: Introdução e Justificativa do Empreendimento.")
    pag2 = doc.new_page()
    pag2.insert_text((50, 50), "Capítulo 2: Sondagem a trado indicando lençol freático a 2,50 m.")
    conteudo = doc.tobytes()
    doc.close()

    texto, info = LeitorPDF.extrair(conteudo)
    assert "[[pag 1]]" in texto
    assert "[[pag 2]]" in texto
    assert info["paginas"] == 2
    assert info["metodo"] == "texto_nativo"

    # Localização de página precisa
    pag_lencol = LeitorPDF.localizar_pagina(texto, "lençol freático")
    assert pag_lencol == 2
    pag_intro = LeitorPDF.localizar_pagina(texto, "Introdução")
    assert pag_intro == 1

    # Fatiamento cirúrgico de páginas (ler página 2 apenas)
    trecho_p2 = LeitorPDF.extrair_paginas(texto, 2, 2)
    assert "lençol freático" in trecho_p2
    assert "Introdução" not in trecho_p2


def test_leitor_pdf_deteccao_texto_corrompido():
    """GabeBrain 07: detecta PDF com camada de texto corrompida / mojibake."""
    from licenciamento.leitor_pdf import LeitorPDF

    # Texto normal
    assert LeitorPDF._texto_corrompido("Este é um documento de licenciamento ambiental perfeitamente legível.") is False

    # Texto corrompido com caracteres nulos / de substituição
    corrompido_subst = "abc\x00\x00\x00\x00\ufffd\ufffd\ufffd\ufffd\ufffd\ufffd\ufffd\ufffd\ufffd" * 5
    assert LeitorPDF._texto_corrompido(corrompido_subst) is True

    # Texto com baixíssima proporção de caracteres alfanuméricos (glifos quebrados)
    corrompido_glifos = "!@#$%^&*()_+{}|:<>?~`!@#$%^&*()_+{}|:<>?~`!@#$%^&*()_+{}|:<>?~`" * 3
    assert LeitorPDF._texto_corrompido(corrompido_glifos) is True


def test_leitor_pdf_estrutura_toc_e_outline_util():
    """GabeBrain 17 (biblioteca-mapa-documento): extração de esqueleto ToC gratuito,
    validação de outline e marcação de seções quentes."""
    import pymupdf
    from licenciamento.leitor_pdf import LeitorPDF

    doc = pymupdf.open()
    for _ in range(10):
        doc.new_page()

    # Sumário com capítulos distintos e relevantes
    toc = [
        [1, "Introdução Geral", 1],
        [1, "Diagnóstico de Sondagem e Lençol Freático", 3],
        [1, "Inventário de Fauna Silvestre", 6],
        [1, "Considerações Finais", 9],
    ]
    doc.set_toc(toc)
    conteudo = doc.tobytes()
    doc.close()

    estrutura = LeitorPDF.extrair_estrutura(conteudo)
    assert estrutura["rota"] == "sumario_embutido"
    assert estrutura["total_paginas"] == 10
    assert len(estrutura["secoes"]) == 4

    # Verifica se detectou seções quentes ambientais
    secoes_quentes = [s for s in estrutura["secoes"] if s["quente"]]
    titulos_quentes = [s["titulo"] for s in secoes_quentes]
    assert any("Sondagem" in t for t in titulos_quentes)
    assert any("Fauna" in t for t in titulos_quentes)


def test_outline_util_rejeita_sumario_invalido():
    """Valida outline_util rejeitando sumários onde todos os itens apontam para a mesma página."""
    from licenciamento.leitor_pdf import LeitorPDF

    # Sumário falso (todos na página 1)
    toc_invalido = [
        (1, "Capítulo 1", 1),
        (1, "Capítulo 2", 1),
        (1, "Capítulo 3", 1),
        (1, "Capítulo 4", 1),
    ]
    serve, motivo = LeitorPDF.outline_util(toc_invalido, 100)
    assert serve is False
    assert "páginas distintas" in motivo or "caem na mesma página" in motivo


def test_auditor_tecnico_rastreamento_pagina():
    """Verifica que o AuditorTecnico inclui rastreamento de página [[pag N]] nos achados."""
    from licenciamento.auditor_tecnico import AuditorTecnico
    from licenciamento.esquemas_tecnicos import StatusValidacao

    at = AuditorTecnico()
    texto_com_paginas = (
        "[[pag 1]]\n"
        "RELATÓRIO TÉCNICO DE MEIO FÍSICO\n"
        "Empreendimento Residencial Vale Verde.\n\n"
        "[[pag 2]]\n"
        "Resultados das investigações geotécnicas:\n"
        "Área do aterro: 1.0 ha.\n"
        "Foram realizados 3 furos de sondagem a trado.\n"
        "Profundidade do lençol freático: 2.0 m.\n"
        "Cota base do aterro: 1.0 m.\n"
        "Distância vertical informada: 1.0 m.\n"
        "Foram executados 2 ensaios de permeabilidade.\n"
    )

    metricas = at.extrair_parametros_sondagem(texto_com_paginas, contexto="RSCC")
    assert metricas.pagina_referencia == 2

    resultado = at.validar_sondagem_aterramento("laudo_geologico.pdf", metricas)
    assert resultado.status == StatusValidacao.PENDENTE
    assert resultado.pagina_referencia == 2
    assert "[pág. 2]" in resultado.trecho_referencia


